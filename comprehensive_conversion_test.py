#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
全面的文档转换测试脚本
测试 Excel、Word、PDF 转 Markdown 的功能
"""

import os
import sys
import traceback
from pathlib import Path

# 添加项目根目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from converter.document_converter import DocumentConverter
from utils.config_manager import ConfigManager

def create_test_excel():
    """创建测试Excel文件"""
    try:
        import pandas as pd
        
        data = {
            '姓名': ['张三', '李四', '王五', '赵六'],
            '年龄': [25, 30, 28, 35],
            '部门': ['技术部', '销售部', '人事部', '财务部'],
            '薪资': [8000, 6000, 5500, 7000]
        }
        
        df = pd.DataFrame(data)
        excel_file = 'test_excel.xlsx'
        df.to_excel(excel_file, index=False, sheet_name='员工信息')
        print(f"✓ 创建测试Excel文件: {excel_file}")
        return excel_file
    except Exception as e:
        print(f"✗ 创建Excel文件失败: {e}")
        return None

def create_test_word():
    """创建测试Word文件"""
    try:
        from docx import Document
        
        doc = Document()
        doc.add_heading('测试文档', 0)
        doc.add_paragraph('这是一个测试Word文档。')
        doc.add_heading('第一章', level=1)
        doc.add_paragraph('这是第一章的内容。包含一些测试文本。')
        
        # 添加表格
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = '项目'
        hdr_cells[1].text = '数量'
        hdr_cells[2].text = '价格'
        
        for i in range(3):
            row_cells = table.add_row().cells
            row_cells[0].text = f'项目{i+1}'
            row_cells[1].text = str(i+1)
            row_cells[2].text = f'{(i+1)*100}元'
        
        word_file = 'test_word.docx'
        doc.save(word_file)
        print(f"✓ 创建测试Word文件: {word_file}")
        return word_file
    except Exception as e:
        print(f"✗ 创建Word文件失败: {e}")
        return None

def create_test_pdf():
    """创建测试PDF文件"""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        
        pdf_file = 'test_pdf.pdf'
        c = canvas.Canvas(pdf_file, pagesize=letter)
        
        # 添加文本内容
        c.drawString(100, 750, "Test PDF Document")
        c.drawString(100, 720, "This is a test PDF file for conversion testing.")
        c.drawString(100, 690, "Chapter 1: Introduction")
        c.drawString(100, 660, "This chapter contains some sample text content.")
        c.drawString(100, 630, "We will test the PDF to Markdown conversion.")
        
        c.save()
        print(f"✓ 创建测试PDF文件: {pdf_file}")
        return pdf_file
    except Exception as e:
        print(f"✗ 创建PDF文件失败: {e}")
        return None

def test_conversion(converter, input_file, target_format):
    """测试单个文件转换"""
    try:
        print(f"\n--- 测试 {input_file} -> {target_format} ---")
        
        # 检查输入文件是否存在
        if not os.path.exists(input_file):
            print(f"✗ 输入文件不存在: {input_file}")
            return False
        
        # 检查是否支持转换
        input_ext = os.path.splitext(input_file)[1].lower()
        if not converter.can_convert(input_ext, target_format):
            print(f"✗ 不支持从 {input_ext} 转换到 {target_format}")
            return False
        
        print(f"✓ 支持转换: {input_ext} -> {target_format}")
        
        # 执行转换
        result = converter.convert_file(input_file, target_format)
        
        if result and os.path.exists(result):
            file_size = os.path.getsize(result)
            print(f"✓ 转换成功: {result} (大小: {file_size} bytes)")
            
            # 显示内容预览
            try:
                with open(result, 'r', encoding='utf-8') as f:
                    content = f.read()
                    lines = content.split('\n')
                    preview_lines = lines[:10]
                    print("内容预览:")
                    for i, line in enumerate(preview_lines, 1):
                        print(f"  {i}: {line[:80]}{'...' if len(line) > 80 else ''}")
                    if len(lines) > 10:
                        print(f"  ... (还有 {len(lines) - 10} 行)")
            except Exception as e:
                print(f"无法读取输出文件内容: {e}")
            
            return True
        else:
            print(f"✗ 转换失败: {result}")
            return False
            
    except Exception as e:
        print(f"✗ 转换过程出错: {e}")
        traceback.print_exc()
        return False

def main():
    """主测试函数"""
    print("=== DocuFlow 全面转换测试 ===\n")
    
    try:
        # 初始化转换器
        print("1. 初始化转换器...")
        config = ConfigManager()
        converter = DocumentConverter(config)
        print("✓ 转换器初始化成功")
        
        # 创建测试文件
        print("\n2. 创建测试文件...")
        test_files = []
        
        excel_file = create_test_excel()
        if excel_file:
            test_files.append(excel_file)
        
        word_file = create_test_word()
        if word_file:
            test_files.append(word_file)
        
        pdf_file = create_test_pdf()
        if pdf_file:
            test_files.append(pdf_file)
        
        if not test_files:
            print("✗ 无法创建任何测试文件")
            return
        
        # 执行转换测试
        print("\n3. 执行转换测试...")
        results = []
        
        for test_file in test_files:
            success = test_conversion(converter, test_file, '.md')
            results.append((test_file, success))
        
        # 显示测试结果
        print("\n=== 测试结果汇总 ===")
        success_count = sum(1 for _, success in results if success)
        total_count = len(results)
        
        for test_file, success in results:
            status = "✓ 成功" if success else "✗ 失败"
            print(f"{test_file}: {status}")
        
        print(f"\n总计: {success_count}/{total_count} 成功")
        print(f"成功率: {success_count/total_count*100:.1f}%")
        
        if success_count == total_count:
            print("\n🎉 所有转换测试通过！")
        else:
            print(f"\n⚠️  有 {total_count - success_count} 个转换失败，需要进一步优化")
        
        # 清理测试文件
        print("\n4. 清理测试文件...")
        for test_file in test_files:
            try:
                if os.path.exists(test_file):
                    os.remove(test_file)
                    print(f"✓ 删除: {test_file}")
                
                # 删除输出文件
                output_file = os.path.splitext(test_file)[0] + '.md'
                if os.path.exists(output_file):
                    os.remove(output_file)
                    print(f"✓ 删除: {output_file}")
            except Exception as e:
                print(f"✗ 删除文件失败 {test_file}: {e}")
        
    except Exception as e:
        print(f"✗ 测试过程出错: {e}")
        traceback.print_exc()

if __name__ == "__main__":
    main()