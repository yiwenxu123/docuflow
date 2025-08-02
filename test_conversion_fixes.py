#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
转换问题修复验证脚本
测试Word转HTML中文编码和HTML转Markdown文件输出问题的修复效果
"""

import os
import sys
import tempfile
from pathlib import Path
from docx import Document
from converter.document_converter import DocumentConverter
from utils.config_manager import ConfigManager

def create_test_word_file(file_path: str) -> str:
    """创建包含中文内容的测试Word文件"""
    doc = Document()
    
    # 添加标题
    doc.add_heading('中文测试文档', 0)
    
    # 添加段落
    doc.add_paragraph('这是一个包含中文内容的测试文档。')
    doc.add_paragraph('测试内容包括：')
    
    # 添加列表
    doc.add_paragraph('• 中文字符编码测试', style='List Bullet')
    doc.add_paragraph('• 特殊符号：￥、©、®、™', style='List Bullet')
    doc.add_paragraph('• 数字和英文：123 ABC', style='List Bullet')
    
    # 添加表格
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '姓名'
    table.cell(0, 1).text = '年龄'
    table.cell(1, 0).text = '张三'
    table.cell(1, 1).text = '25'
    
    # 保存文件
    doc.save(file_path)
    print(f"✓ 创建测试Word文件: {file_path}")
    return file_path

def create_test_html_file(file_path: str) -> str:
    """创建包含中文内容的测试HTML文件"""
    html_content = """
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>中文HTML测试</title>
</head>
<body>
    <h1>中文HTML测试文档</h1>
    <p>这是一个包含中文内容的HTML文档。</p>
    
    <h2>测试内容</h2>
    <ul>
        <li>中文字符编码测试</li>
        <li>特殊符号：￥、©、®、™</li>
        <li>数字和英文：123 ABC</li>
    </ul>
    
    <table border="1">
        <tr>
            <th>姓名</th>
            <th>年龄</th>
        </tr>
        <tr>
            <td>张三</td>
            <td>25</td>
        </tr>
    </table>
    
    <p><strong>粗体文本</strong> 和 <em>斜体文本</em></p>
    <p><a href="#">链接测试</a></p>
</body>
</html>
    """
    
    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print(f"✓ 创建测试HTML文件: {file_path}")
    return file_path

def test_word_to_html_encoding(converter: DocumentConverter, test_dir: str) -> bool:
    """测试Word转HTML的中文编码"""
    print("\n=== 测试Word转HTML中文编码 ===")
    
    try:
        # 创建测试文件
        word_file = os.path.join(test_dir, "test_chinese.docx")
        html_file = os.path.join(test_dir, "test_chinese.html")
        
        create_test_word_file(word_file)
        
        # 执行转换
        print("执行Word转HTML转换...")
        result_path = converter.convert_file(word_file, html_file)
        
        # 验证输出文件
        if not os.path.exists(result_path):
            print(f"❌ 输出文件不存在: {result_path}")
            return False
        
        # 检查文件内容编码
        with open(result_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 验证中文内容
        chinese_keywords = ['中文测试文档', '这是一个包含中文内容', '张三', '姓名', '年龄']
        missing_keywords = []
        
        for keyword in chinese_keywords:
            if keyword not in content:
                missing_keywords.append(keyword)
        
        if missing_keywords:
            print(f"❌ HTML文件中缺少中文内容: {missing_keywords}")
            print(f"文件内容预览: {content[:200]}...")
            return False
        
        # 检查HTML编码声明
        if 'charset=utf-8' not in content.lower() and 'charset="utf-8"' not in content.lower():
            print("⚠️  HTML文件中未找到UTF-8编码声明")
        
        print(f"✓ Word转HTML中文编码测试通过")
        print(f"  输出文件: {result_path}")
        print(f"  文件大小: {os.path.getsize(result_path)} 字节")
        return True
        
    except Exception as e:
        print(f"❌ Word转HTML测试失败: {str(e)}")
        return False

def test_html_to_markdown_output(converter: DocumentConverter, test_dir: str) -> bool:
    """测试HTML转Markdown的文件输出"""
    print("\n=== 测试HTML转Markdown文件输出 ===")
    
    try:
        # 创建测试文件
        html_file = os.path.join(test_dir, "test_chinese.html")
        md_file = os.path.join(test_dir, "test_chinese.md")
        
        create_test_html_file(html_file)
        
        # 执行转换
        print("执行HTML转Markdown转换...")
        result_path = converter.convert_file(html_file, md_file)
        
        # 验证输出文件
        if not os.path.exists(result_path):
            print(f"❌ 输出文件不存在: {result_path}")
            return False
        
        # 检查文件大小
        file_size = os.path.getsize(result_path)
        if file_size == 0:
            print(f"❌ 输出文件为空: {result_path}")
            return False
        
        # 检查文件内容
        with open(result_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 验证Markdown格式和中文内容
        markdown_indicators = ['#', '*', '-', '|']
        chinese_keywords = ['中文HTML测试文档', '这是一个包含中文内容', '张三', '姓名', '年龄']
        
        has_markdown_format = any(indicator in content for indicator in markdown_indicators)
        missing_keywords = [keyword for keyword in chinese_keywords if keyword not in content]
        
        if not has_markdown_format:
            print("⚠️  输出文件可能不是有效的Markdown格式")
        
        if missing_keywords:
            print(f"❌ Markdown文件中缺少中文内容: {missing_keywords}")
            print(f"文件内容预览: {content[:200]}...")
            return False
        
        print(f"✓ HTML转Markdown文件输出测试通过")
        print(f"  输出文件: {result_path}")
        print(f"  文件大小: {file_size} 字节")
        print(f"  内容预览: {content[:100]}...")
        return True
        
    except Exception as e:
        print(f"❌ HTML转Markdown测试失败: {str(e)}")
        return False

def test_end_to_end_conversion(converter: DocumentConverter, test_dir: str) -> bool:
    """测试Word -> HTML -> Markdown的完整转换流程"""
    print("\n=== 测试完整转换流程 (Word -> HTML -> Markdown) ===")
    
    try:
        # 创建测试文件
        word_file = os.path.join(test_dir, "test_flow.docx")
        html_file = os.path.join(test_dir, "test_flow.html")
        md_file = os.path.join(test_dir, "test_flow.md")
        
        create_test_word_file(word_file)
        
        # 第一步：Word转HTML
        print("步骤1: Word转HTML...")
        html_result = converter.convert_file(word_file, html_file)
        
        if not os.path.exists(html_result):
            print(f"❌ HTML文件生成失败: {html_result}")
            return False
        
        # 第二步：HTML转Markdown
        print("步骤2: HTML转Markdown...")
        md_result = converter.convert_file(html_result, md_file)
        
        if not os.path.exists(md_result):
            print(f"❌ Markdown文件生成失败: {md_result}")
            return False
        
        # 验证最终结果
        with open(md_result, 'r', encoding='utf-8') as f:
            final_content = f.read()
        
        chinese_keywords = ['中文测试文档', '张三']
        missing_keywords = [keyword for keyword in chinese_keywords if keyword not in final_content]
        
        if missing_keywords:
            print(f"❌ 最终Markdown文件中缺少关键内容: {missing_keywords}")
            return False
        
        print(f"✓ 完整转换流程测试通过")
        print(f"  Word文件: {word_file}")
        print(f"  HTML文件: {html_result} ({os.path.getsize(html_result)} 字节)")
        print(f"  Markdown文件: {md_result} ({os.path.getsize(md_result)} 字节)")
        return True
        
    except Exception as e:
        print(f"❌ 完整转换流程测试失败: {str(e)}")
        return False

def cleanup_test_files(test_dir: str):
    """清理测试文件"""
    try:
        import shutil
        if os.path.exists(test_dir):
            shutil.rmtree(test_dir)
            print(f"\n✓ 清理测试文件: {test_dir}")
    except Exception as e:
        print(f"⚠️  清理测试文件失败: {str(e)}")

def main():
    """主测试函数"""
    print("转换问题修复验证测试")
    print("=" * 50)
    
    # 创建临时测试目录
    test_dir = tempfile.mkdtemp(prefix="conversion_fix_test_")
    print(f"测试目录: {test_dir}")
    
    try:
        # 初始化转换器
        config_manager = ConfigManager()
        config = config_manager.get_config()
        converter = DocumentConverter(config_manager)
        
        # 运行测试
        tests = [
            ("Word转HTML中文编码", lambda: test_word_to_html_encoding(converter, test_dir)),
            ("HTML转Markdown文件输出", lambda: test_html_to_markdown_output(converter, test_dir)),
            ("完整转换流程", lambda: test_end_to_end_conversion(converter, test_dir))
        ]
        
        results = []
        for test_name, test_func in tests:
            try:
                result = test_func()
                results.append((test_name, result))
            except Exception as e:
                print(f"❌ {test_name}测试异常: {str(e)}")
                results.append((test_name, False))
        
        # 汇总结果
        print("\n" + "=" * 50)
        print("测试结果汇总:")
        
        passed = 0
        for test_name, result in results:
            status = "✓ 通过" if result else "❌ 失败"
            print(f"  {test_name}: {status}")
            if result:
                passed += 1
        
        print(f"\n总计: {passed}/{len(results)} 项测试通过")
        
        if passed == len(results):
            print("\n🎉 所有转换问题已修复！")
            print("\n修复内容:")
            print("1. Word转HTML: 添加UTF-8编码支持，解决中文乱码问题")
            print("2. HTML转Markdown: 增强文件输出验证和错误处理")
            print("3. 改进编码处理和路径管理")
        else:
            print("\n⚠️  部分测试未通过，请检查相关问题")
        
    finally:
        # 清理测试文件
        cleanup_test_files(test_dir)

if __name__ == "__main__":
    main()