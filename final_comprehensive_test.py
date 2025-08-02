#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DocuFlow 最终综合测试
确保所有转换功能稳定可靠
"""

import os
import sys
import traceback
import time

# 添加项目根目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from converter.document_converter import DocumentConverter
from utils.config_manager import ConfigManager

def create_comprehensive_test_files():
    """创建全面的测试文件"""
    print("=== 创建测试文件 ===")
    test_files = {}
    
    # 创建复杂Excel文件
    try:
        import pandas as pd
        
        # 创建多个工作表的Excel文件
        with pd.ExcelWriter('comprehensive_test.xlsx', engine='openpyxl') as writer:
            # 员工信息表
            employees = pd.DataFrame({
                '姓名': ['张三', '李四', '王五', '赵六', '钱七'],
                '年龄': [25, 30, 28, 35, 32],
                '部门': ['技术部', '销售部', '人事部', '财务部', '市场部'],
                '薪资': [8000, 6000, 5500, 7000, 6500],
                '入职日期': ['2020-01-15', '2019-03-20', '2021-06-10', '2018-11-05', '2020-09-12']
            })
            employees.to_excel(writer, sheet_name='员工信息', index=False)
            
            # 销售数据表
            sales = pd.DataFrame({
                '月份': ['2023-01', '2023-02', '2023-03', '2023-04', '2023-05'],
                '销售额': [120000, 135000, 98000, 156000, 142000],
                '成本': [80000, 90000, 65000, 104000, 95000],
                '利润': [40000, 45000, 33000, 52000, 47000]
            })
            sales.to_excel(writer, sheet_name='销售数据', index=False)
        
        test_files['excel'] = 'comprehensive_test.xlsx'
        print(f"✓ 创建复杂Excel文件: comprehensive_test.xlsx")
    except Exception as e:
        print(f"✗ 创建Excel文件失败: {e}")
    
    # 创建复杂Word文件
    try:
        from docx import Document
        from docx.shared import Inches
        
        doc = Document()
        
        # 添加标题
        doc.add_heading('DocuFlow 测试文档', 0)
        
        # 添加段落
        doc.add_paragraph('这是一个用于测试DocuFlow转换功能的综合性文档。')
        
        # 添加章节
        doc.add_heading('第一章：系统概述', level=1)
        doc.add_paragraph('DocuFlow是一个强大的文档转换工具，支持多种格式之间的转换。')
        
        doc.add_heading('第二章：功能特性', level=1)
        doc.add_paragraph('主要功能包括：')
        
        # 添加列表
        doc.add_paragraph('• Excel转换功能', style='List Bullet')
        doc.add_paragraph('• Word转换功能', style='List Bullet')
        doc.add_paragraph('• PDF转换功能', style='List Bullet')
        
        # 添加表格
        doc.add_heading('第三章：性能数据', level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = '格式'
        hdr_cells[1].text = '转换速度'
        hdr_cells[2].text = '准确率'
        hdr_cells[3].text = '支持度'
        
        data = [
            ['Excel', '快速', '95%', '完全支持'],
            ['Word', '快速', '98%', '完全支持'],
            ['PDF', '中等', '90%', '完全支持']
        ]
        
        for row_data in data:
            row_cells = table.add_row().cells
            for i, cell_data in enumerate(row_data):
                row_cells[i].text = cell_data
        
        doc.save('comprehensive_test.docx')
        test_files['word'] = 'comprehensive_test.docx'
        print(f"✓ 创建复杂Word文件: comprehensive_test.docx")
    except Exception as e:
        print(f"✗ 创建Word文件失败: {e}")
    
    # 创建复杂PDF文件
    try:
        from fpdf import FPDF
        
        class PDF(FPDF):
            def header(self):
                self.set_font('Arial', 'B', 15)
                self.cell(0, 10, 'DocuFlow Test Document', 0, 1, 'C')
                self.ln(10)
            
            def footer(self):
                self.set_y(-15)
                self.set_font('Arial', 'I', 8)
                self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')
        
        pdf = PDF()
        pdf.add_page()
        
        # 添加内容
        pdf.set_font('Arial', 'B', 16)
        pdf.cell(0, 10, 'Chapter 1: Introduction', 0, 1)
        pdf.ln(5)
        
        pdf.set_font('Arial', '', 12)
        pdf.multi_cell(0, 10, 'This is a comprehensive test document for DocuFlow conversion capabilities. '
                              'It contains various types of content including headers, paragraphs, and structured data.')
        pdf.ln(10)
        
        pdf.set_font('Arial', 'B', 14)
        pdf.cell(0, 10, 'Chapter 2: Features', 0, 1)
        pdf.ln(5)
        
        pdf.set_font('Arial', '', 12)
        features = [
            '1. Excel to Markdown conversion',
            '2. Word to Markdown conversion', 
            '3. PDF to Markdown conversion',
            '4. Multi-format support',
            '5. High accuracy processing'
        ]
        
        for feature in features:
            pdf.cell(0, 8, feature, 0, 1)
        
        pdf.output('comprehensive_test.pdf')
        test_files['pdf'] = 'comprehensive_test.pdf'
        print(f"✓ 创建复杂PDF文件: comprehensive_test.pdf")
    except Exception as e:
        print(f"✗ 创建PDF文件失败: {e}")
    
    return test_files

def test_all_conversions(converter, test_files):
    """测试所有转换组合"""
    print("\n=== 全面转换测试 ===")
    
    target_formats = ['.md', '.html', '.txt']
    results = {}
    
    for file_type, file_path in test_files.items():
        print(f"\n--- 测试 {file_type.upper()} 文件转换 ---")
        
        if not os.path.exists(file_path):
            print(f"✗ 文件不存在: {file_path}")
            continue
        
        file_size = os.path.getsize(file_path)
        print(f"输入文件: {file_path} (大小: {file_size} bytes)")
        
        for target_format in target_formats:
            test_key = f"{file_type}_{target_format[1:]}"
            
            try:
                start_time = time.time()
                result = converter.convert_file(file_path, target_format)
                end_time = time.time()
                
                if result and os.path.exists(result):
                    output_size = os.path.getsize(result)
                    duration = end_time - start_time
                    
                    print(f"  ✓ {file_type} -> {target_format}: {result}")
                    print(f"    大小: {output_size} bytes, 耗时: {duration:.2f}s")
                    
                    # 检查内容质量
                    with open(result, 'r', encoding='utf-8') as f:
                        content = f.read()
                        if content.strip():
                            lines = len(content.split('\n'))
                            print(f"    内容: {len(content)} 字符, {lines} 行")
                            results[test_key] = True
                        else:
                            print(f"    ⚠️  内容为空")
                            results[test_key] = False
                else:
                    print(f"  ✗ {file_type} -> {target_format}: 转换失败")
                    results[test_key] = False
                    
            except Exception as e:
                print(f"  ✗ {file_type} -> {target_format}: 出错 - {e}")
                results[test_key] = False
    
    return results

def analyze_results(results):
    """分析测试结果"""
    print("\n=== 测试结果分析 ===")
    
    total_tests = len(results)
    successful_tests = sum(1 for success in results.values() if success)
    failed_tests = total_tests - successful_tests
    
    print(f"总测试数: {total_tests}")
    print(f"成功: {successful_tests}")
    print(f"失败: {failed_tests}")
    print(f"成功率: {successful_tests/total_tests*100:.1f}%")
    
    if failed_tests > 0:
        print("\n失败的测试:")
        for test_name, success in results.items():
            if not success:
                print(f"  ✗ {test_name}")
    
    # 按格式分析
    print("\n按格式分析:")
    formats = ['excel', 'word', 'pdf']
    targets = ['md', 'html', 'txt']
    
    for fmt in formats:
        fmt_results = [results.get(f"{fmt}_{target}", False) for target in targets]
        fmt_success = sum(fmt_results)
        print(f"  {fmt.upper()}: {fmt_success}/{len(targets)} 成功")
    
    # 按目标格式分析
    print("\n按目标格式分析:")
    for target in targets:
        target_results = [results.get(f"{fmt}_{target}", False) for fmt in formats]
        target_success = sum(target_results)
        print(f"  {target.upper()}: {target_success}/{len(formats)} 成功")

def cleanup_files(test_files):
    """清理所有测试文件"""
    print("\n=== 清理测试文件 ===")
    
    # 收集所有可能的文件
    all_files = list(test_files.values())
    
    # 添加输出文件
    for file_path in test_files.values():
        base_name = os.path.splitext(file_path)[0]
        for ext in ['.md', '.html', '.txt']:
            all_files.append(f"{base_name}{ext}")
    
    cleaned_count = 0
    for file_path in all_files:
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                print(f"✓ 删除: {file_path}")
                cleaned_count += 1
        except Exception as e:
            print(f"✗ 删除失败 {file_path}: {e}")
    
    print(f"共清理 {cleaned_count} 个文件")

def main():
    """主函数"""
    print("DocuFlow 最终综合测试")
    print("=" * 60)
    
    try:
        # 初始化转换器
        print("初始化转换器...")
        config = ConfigManager()
        converter = DocumentConverter(config)
        print("✓ 转换器初始化成功")
        
        # 创建测试文件
        test_files = create_comprehensive_test_files()
        if not test_files:
            print("❌ 无法创建测试文件")
            return
        
        # 执行全面测试
        results = test_all_conversions(converter, test_files)
        
        # 分析结果
        analyze_results(results)
        
        # 清理文件
        cleanup_files(test_files)
        
        # 最终结论
        success_rate = sum(1 for success in results.values() if success) / len(results) * 100
        print("\n" + "=" * 60)
        if success_rate >= 90:
            print("🎉 测试完成！系统运行优秀！")
        elif success_rate >= 70:
            print("✅ 测试完成！系统运行良好！")
        else:
            print("⚠️  测试完成！系统需要进一步优化！")
        
        print(f"最终成功率: {success_rate:.1f}%")
        
    except Exception as e:
        print(f"❌ 测试过程出错: {e}")
        traceback.print_exc()

if __name__ == "__main__":
    main()