#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
全面测试Web版本所有功能
"""

import requests
import json
import time
import os
from pathlib import Path
from docx import Document
import pandas as pd

def create_test_files():
    """创建测试文件"""
    print(f"📁 创建测试文件... 当前工作目录: {os.getcwd()}")
    
    # 创建Word文档
    doc = Document()
    doc.add_heading('测试文档', 0)
    doc.add_paragraph('这是一个测试Word文档，用于验证转换功能。')
    doc.add_paragraph('包含多种格式：')
    doc.add_paragraph('• 列表项目1')
    doc.add_paragraph('• 列表项目2')
    doc.add_heading('子标题', level=1)
    doc.add_paragraph('这是子标题下的内容。')
    
    word_file = "test_comprehensive.docx"
    doc.save(word_file)
    print(f"✅ Word文件创建: {word_file}")
    
    # 创建Markdown文件
    md_content = """# 测试Markdown文档

这是一个测试Markdown文档。

## 功能列表

- 支持标题
- 支持列表
- 支持**粗体**和*斜体*
- 支持代码块

```python
def hello():
    print("Hello, World!")
```

## 表格

| 列1 | 列2 | 列3 |
|-----|-----|-----|
| 数据1 | 数据2 | 数据3 |
| 数据4 | 数据5 | 数据6 |
"""
    
    md_file = "test_comprehensive.md"
    with open(md_file, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"✅ Markdown文件创建: {md_file}")
    
    # 创建Excel文件
    data = {
        '姓名': ['张三', '李四', '王五'],
        '年龄': [25, 30, 35],
        '城市': ['北京', '上海', '广州']
    }
    df = pd.DataFrame(data)
    excel_file = "test_comprehensive.xlsx"
    df.to_excel(excel_file, index=False)
    print(f"✅ Excel文件创建: {excel_file}")
    
    # 创建HTML文件
    html_content = """<!DOCTYPE html>
<html>
<head>
    <title>测试HTML文档</title>
</head>
<body>
    <h1>测试HTML文档</h1>
    <p>这是一个测试HTML文档。</p>
    <ul>
        <li>列表项目1</li>
        <li>列表项目2</li>
    </ul>
    <table border="1">
        <tr>
            <th>列1</th>
            <th>列2</th>
        </tr>
        <tr>
            <td>数据1</td>
            <td>数据2</td>
        </tr>
    </table>
</body>
</html>"""
    
    html_file = "test_comprehensive.html"
    with open(html_file, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"✅ HTML文件创建: {html_file}")
    
    return [word_file, md_file, excel_file, html_file]

def test_server_connection():
    """测试服务器连接"""
    print("🔗 测试服务器连接...")
    try:
        response = requests.get('http://localhost:5003')
        if response.status_code == 200:
            print("✅ 服务器连接正常")
            return True
        else:
            print(f"❌ 服务器响应异常: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 服务器连接失败: {e}")
        return False

def test_file_upload(file_path):
    """测试文件上传"""
    print(f"📤 测试文件上传: {file_path}")
    try:
        with open(file_path, 'rb') as f:
            files = {'files': (file_path, f, 'application/octet-stream')}
            response = requests.post('http://localhost:5003/upload', files=files)
        
        if response.status_code == 200:
            data = response.json()
            if data.get('success'):
                print(f"✅ 文件上传成功: {data.get('message')}")
                return data.get('files', [])
            else:
                print(f"❌ 文件上传失败: {data.get('error')}")
                return None
        else:
            print(f"❌ 上传请求失败: {response.status_code}")
            return None
    except Exception as e:
        print(f"❌ 上传异常: {e}")
        return None

def test_conversion(files, output_format):
    """测试文件转换"""
    print(f"🔄 测试转换到 {output_format} 格式...")
    try:
        conversion_data = {
            'files': [f['path'] for f in files],
            'output_format': output_format,
            'keep_original_name': True,
            'use_async': False,
            'max_workers': 2,
            'output_location': 'temp'
        }
        
        response = requests.post(
            'http://localhost:5003/convert',
            headers={'Content-Type': 'application/json'},
            data=json.dumps(conversion_data)
        )
        
        if response.status_code == 200:
            data = response.json()
            if data.get('success'):
                summary = data.get('summary', {})
                print(f"✅ 转换完成: {summary.get('success', 0)}/{summary.get('total', 0)} 成功")
                print(f"📈 成功率: {summary.get('success_rate', 0):.1f}%")
                
                # 显示转换结果
                for result in data.get('results', []):
                    if result.get('success'):
                        print(f"  ✅ {result.get('file')} -> 下载ID: {result.get('download_id')}")
                    else:
                        print(f"  ❌ {result.get('file')} -> 错误: {result.get('error')}")
                
                return data
            else:
                print(f"❌ 转换失败: {data.get('error')}")
                return None
        else:
            print(f"❌ 转换请求失败: {response.status_code}")
            return None
    except Exception as e:
        print(f"❌ 转换异常: {e}")
        return None

def test_stats():
    """测试统计功能"""
    print("📊 测试统计功能...")
    try:
        response = requests.get('http://localhost:5003/stats')
        if response.status_code == 200:
            data = response.json()
            if data.get('success'):
                stats = data.get('stats', {})
                print(f"✅ 统计信息获取成功:")
                print(f"  总转换次数: {stats.get('total_conversions', 0)}")
                print(f"  成功转换: {stats.get('successful_conversions', 0)}")
                print(f"  成功率: {stats.get('success_rate', 0):.1f}%")
                print(f"  平均转换时间: {stats.get('average_conversion_time', 0):.2f}s")
                return True
            else:
                print(f"❌ 统计获取失败: {data.get('error')}")
                return False
        else:
            print(f"❌ 统计请求失败: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 统计异常: {e}")
        return False

def test_cleanup():
    """测试清理功能"""
    print("🧹 测试清理功能...")
    try:
        response = requests.post('http://localhost:5003/cleanup')
        if response.status_code == 200:
            data = response.json()
            if data.get('success'):
                print(f"✅ 清理成功: {data.get('message')}")
                return True
            else:
                print(f"❌ 清理失败: {data.get('error')}")
                return False
        else:
            print(f"❌ 清理请求失败: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 清理异常: {e}")
        return False

def main():
    """主测试函数"""
    print("🧪 开始全面测试Web版本功能...")
    print("=" * 50)
    
    # 1. 测试服务器连接
    if not test_server_connection():
        print("❌ 服务器连接失败，测试终止")
        return
    
    # 2. 创建测试文件
    test_files = create_test_files()
    
    # 3. 测试各种转换组合
    conversion_tests = [
        ('.md', ['test_comprehensive.docx', 'test_comprehensive.html']),
        ('.html', ['test_comprehensive.md', 'test_comprehensive.docx']),
        ('.docx', ['test_comprehensive.md', 'test_comprehensive.html']),
    ]
    
    success_count = 0
    total_tests = 0
    
    for output_format, input_files in conversion_tests:
        print(f"\n📋 测试转换到 {output_format} 格式")
        print("-" * 30)
        
        for input_file in input_files:
            if os.path.exists(input_file):
                total_tests += 1
                print(f"\n🔄 {input_file} -> {output_format}")
                
                # 上传文件
                uploaded_files = test_file_upload(input_file)
                if uploaded_files:
                    # 转换文件
                    result = test_conversion(uploaded_files, output_format)
                    if result and result.get('success'):
                        success_count += 1
                        print("✅ 转换测试通过")
                    else:
                        print("❌ 转换测试失败")
                else:
                    print("❌ 上传测试失败")
    
    # 4. 测试其他功能
    print(f"\n📊 测试其他功能")
    print("-" * 30)
    
    stats_ok = test_stats()
    cleanup_ok = test_cleanup()
    
    # 5. 清理测试文件
    print(f"\n🧹 清理测试文件...")
    for file in test_files:
        try:
            if os.path.exists(file):
                os.remove(file)
                print(f"✅ 删除: {file}")
        except Exception as e:
            print(f"❌ 删除失败 {file}: {e}")
    
    # 6. 测试总结
    print(f"\n📋 测试总结")
    print("=" * 50)
    print(f"转换测试: {success_count}/{total_tests} 通过")
    print(f"统计功能: {'✅ 正常' if stats_ok else '❌ 异常'}")
    print(f"清理功能: {'✅ 正常' if cleanup_ok else '❌ 异常'}")
    
    overall_success = (success_count == total_tests) and stats_ok and cleanup_ok
    print(f"\n🎯 总体结果: {'✅ 所有功能正常' if overall_success else '❌ 部分功能异常'}")
    
    if overall_success:
        print("\n🎉 Web版本功能测试全部通过！")
    else:
        print("\n⚠️ 发现问题，请检查上述错误信息")

if __name__ == "__main__":
    main()