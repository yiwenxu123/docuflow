#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
测试Word转Markdown功能
"""

import requests
import os
import tempfile
from docx import Document

def create_test_word_file():
    """创建测试Word文件"""
    doc = Document()
    doc.add_heading('测试文档标题', 0)
    doc.add_heading('一级标题', level=1)
    doc.add_paragraph('这是一个测试段落，包含一些基本文本内容。')
    
    doc.add_heading('二级标题', level=2)
    doc.add_paragraph('这里是另一个段落，用于测试Word到Markdown的转换功能。')
    
    # 添加列表
    doc.add_paragraph('项目1', style='List Bullet')
    doc.add_paragraph('项目2', style='List Bullet')
    doc.add_paragraph('项目3', style='List Bullet')
    
    # 保存文件
    test_file = "/Users/yiwenxu123/Documents/tare/文档转换/test_word_conversion.docx"
    doc.save(test_file)
    return test_file

def test_word_to_markdown():
    """测试Word转Markdown功能"""
    base_url = "http://localhost:5003"
    
    print("🧪 开始测试Word转Markdown功能...")
    
    # 1. 创建测试Word文件
    print("1️⃣ 创建测试Word文件...")
    test_file = create_test_word_file()
    print(f"✅ 测试文件创建成功: {test_file}")
    
    # 2. 测试服务器连接
    print("2️⃣ 测试服务器连接...")
    try:
        response = requests.get(f"{base_url}/")
        if response.status_code == 200:
            print("✅ 服务器连接正常")
        else:
            print(f"❌ 服务器连接失败: {response.status_code}")
            return
    except Exception as e:
        print(f"❌ 服务器连接异常: {e}")
        return
    
    # 3. 上传文件
    print("3️⃣ 测试文件上传...")
    try:
        with open(test_file, 'rb') as f:
            files = {'files': (os.path.basename(test_file), f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
            response = requests.post(f"{base_url}/upload", files=files)
        
        if response.status_code == 200:
            upload_result = response.json()
            print("✅ 文件上传成功")
            print(f"📄 上传文件数: {upload_result.get('file_count', 0)}")
            if 'files' in upload_result:
                for file_info in upload_result['files']:
                    print(f"📝 文件信息: {file_info['name']} ({file_info['size']}MB)")
        else:
            print(f"❌ 文件上传失败: {response.status_code}")
            print(f"错误信息: {response.text}")
            return
    except Exception as e:
        print(f"❌ 文件上传异常: {e}")
        return
    
    # 4. 转换文件
    print("4️⃣ 测试文件转换 (DOCX -> Markdown)...")
    try:
        conversion_data = {
            'files': [os.path.basename(test_file)],
            'output_format': 'md'
        }
        
        response = requests.post(f"{base_url}/convert", json=conversion_data)
        
        if response.status_code == 200:
            result = response.json()
            print("✅ 文件转换成功")
            
            # 使用正确的响应格式
            summary = result.get('summary', {})
            print(f"📊 转换统计: {summary.get('success', 0)}/{summary.get('total', 0)} 成功")
            print(f"📈 成功率: {summary.get('success_rate', 0):.1f}%")
            
            # 显示转换结果
            if 'results' in result:
                for res in result['results']:
                    if res['success']:
                        print(f"✅ 输出文件: {os.path.basename(res['output_file'])}")
                    else:
                        print(f"❌ 转换失败: {res['error']}")
        else:
            print(f"❌ 文件转换失败: {response.status_code}")
            print(f"错误信息: {response.text}")
            return
    except Exception as e:
        print(f"❌ 文件转换异常: {e}")
        return
    
    print("\n" + "="*50)
    print("🎉 Word转Markdown功能测试完成!")

if __name__ == "__main__":
    test_word_to_markdown()