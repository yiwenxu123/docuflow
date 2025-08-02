#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Web版本上传功能调试测试
"""

import requests
import json
import time
import os
from pathlib import Path
from docx import Document

def create_test_file():
    """创建测试文件"""
    print("📁 创建测试文件...")
    
    # 创建Word文档
    doc = Document()
    doc.add_heading('测试文档', 0)
    doc.add_paragraph('这是一个测试Word文档，用于验证Web上传功能。')
    doc.add_paragraph('包含中文内容和特殊字符：！@#￥%……&*（）')
    
    word_file = "web_test.docx"
    doc.save(word_file)
    print(f"✅ 测试文件创建: {word_file}")
    return word_file

def test_server_connection():
    """测试服务器连接"""
    try:
        response = requests.get('http://localhost:5003/', timeout=5)
        if response.status_code == 200:
            print("✅ 服务器连接正常")
            return True
        else:
            print(f"❌ 服务器响应异常: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 服务器连接失败: {e}")
        return False

def test_file_upload(file_path):
    """测试文件上传"""
    print(f"\n📤 测试文件上传: {file_path}")
    
    if not os.path.exists(file_path):
        print(f"❌ 文件不存在: {file_path}")
        return False
    
    try:
        with open(file_path, 'rb') as f:
            files = {'files': (os.path.basename(file_path), f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
            response = requests.post('http://localhost:5003/upload', files=files, timeout=30)
        
        print(f"📊 响应状态码: {response.status_code}")
        print(f"📊 响应头: {dict(response.headers)}")
        
        if response.status_code == 200:
            result = response.json()
            print(f"📊 响应内容: {json.dumps(result, ensure_ascii=False, indent=2)}")
            
            if result.get('success'):
                print("✅ 文件上传成功")
                return result.get('files', [])
            else:
                print(f"❌ 上传失败: {result.get('error', '未知错误')}")
                return False
        else:
            print(f"❌ HTTP错误: {response.status_code}")
            print(f"响应内容: {response.text}")
            return False
            
    except Exception as e:
        print(f"❌ 上传异常: {e}")
        return False

def test_conversion(uploaded_files, target_format='.md'):
    """测试文件转换"""
    print(f"\n🔄 测试转换到 {target_format} 格式...")
    
    if not uploaded_files:
        print("❌ 没有已上传的文件")
        return False
    
    file_paths = [file_info['path'] for file_info in uploaded_files]
    
    try:
        data = {
            'files': file_paths,
            'output_format': target_format,
            'keep_original_name': True,
            'use_async': False,
            'output_location': 'temp'
        }
        
        response = requests.post('http://localhost:5003/convert', 
                               json=data, 
                               headers={'Content-Type': 'application/json'},
                               timeout=60)
        
        print(f"📊 转换响应状态码: {response.status_code}")
        
        if response.status_code == 200:
            result = response.json()
            print(f"📊 转换结果: {json.dumps(result, ensure_ascii=False, indent=2)}")
            
            success_count = sum(1 for r in result.get('results', []) if r.get('success'))
            total_count = len(result.get('results', []))
            
            print(f"✅ 转换完成: {success_count}/{total_count} 成功")
            return result.get('results', [])
        else:
            print(f"❌ 转换HTTP错误: {response.status_code}")
            print(f"响应内容: {response.text}")
            return False
            
    except Exception as e:
        print(f"❌ 转换异常: {e}")
        return False

def test_stats():
    """测试统计功能"""
    print(f"\n📊 测试统计功能...")
    
    try:
        response = requests.get('http://localhost:5003/stats', timeout=10)
        
        if response.status_code == 200:
            stats = response.json()
            print(f"✅ 统计信息: {json.dumps(stats, ensure_ascii=False, indent=2)}")
            return True
        else:
            print(f"❌ 统计功能错误: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ 统计功能异常: {e}")
        return False

def cleanup_test_files():
    """清理测试文件"""
    print(f"\n🧹 清理测试文件...")
    
    test_files = ['web_test.docx']
    for file_path in test_files:
        if os.path.exists(file_path):
            os.remove(file_path)
            print(f"✅ 删除: {file_path}")

def main():
    """主测试函数"""
    print("🧪 开始Web版本上传功能调试测试...")
    print("=" * 50)
    
    # 1. 测试服务器连接
    if not test_server_connection():
        print("❌ 服务器连接失败，测试终止")
        return
    
    # 2. 创建测试文件
    test_file = create_test_file()
    
    try:
        # 3. 测试文件上传
        uploaded_files = test_file_upload(test_file)
        if not uploaded_files:
            print("❌ 文件上传失败，测试终止")
            return
        
        # 4. 测试文件转换
        conversion_results = test_conversion(uploaded_files, '.md')
        if not conversion_results:
            print("❌ 文件转换失败")
        
        # 5. 测试统计功能
        test_stats()
        
        print("\n📋 测试总结")
        print("=" * 50)
        if uploaded_files and conversion_results:
            print("🎯 总体结果: ✅ 所有功能正常")
        else:
            print("🎯 总体结果: ❌ 部分功能异常")
            
    finally:
        # 6. 清理测试文件
        cleanup_test_files()

if __name__ == "__main__":
    main()