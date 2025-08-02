#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DocuFlow 转换功能诊断和优化工具
检查并修复 Excel、Word、PDF 转 Markdown 的问题
"""

import os
import sys
import traceback
import subprocess
from pathlib import Path

# 添加项目根目录到路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

def check_dependencies():
    """检查依赖包"""
    print("=== 检查依赖包 ===")
    
    required_packages = {
        'pandas': 'Excel处理',
        'openpyxl': 'Excel读写',
        'docx': 'Word文档处理',
        'pypdf': 'PDF处理',
        'pdfplumber': 'PDF文本提取',
        'fpdf': 'PDF创建',
        'markdown': 'Markdown处理'
    }
    
    missing_packages = []
    
    for package, description in required_packages.items():
        try:
            __import__(package.replace('-', '_'))
            print(f"✓ {package}: {description}")
        except ImportError:
            print(f"✗ {package}: {description} - 缺失")
            missing_packages.append(package)
    
    if missing_packages:
        print(f"\n需要安装的包: {', '.join(missing_packages)}")
        print("运行命令: pip install " + " ".join(missing_packages))
        return False
    else:
        print("\n✓ 所有依赖包都已安装")
        return True

def check_pandoc():
    """检查Pandoc"""
    print("\n=== 检查Pandoc ===")
    try:
        result = subprocess.run(['pandoc', '--version'], 
                              capture_output=True, text=True, timeout=10)
        if result.returncode == 0:
            version_line = result.stdout.split('\n')[0]
            print(f"✓ Pandoc已安装: {version_line}")
            return True
        else:
            print("✗ Pandoc未正确安装")
            return False
    except (subprocess.TimeoutExpired, FileNotFoundError):
        print("✗ Pandoc未安装或不在PATH中")
        print("请安装Pandoc: https://pandoc.org/installing.html")
        return False

def test_converter_initialization():
    """测试转换器初始化"""
    print("\n=== 测试转换器初始化 ===")
    try:
        from converter.document_converter import DocumentConverter
        from utils.config_manager import ConfigManager
        
        config = ConfigManager()
        converter = DocumentConverter(config)
        print("✓ 转换器初始化成功")
        
        # 检查支持的转换
        conversions_to_test = [
            ('.xlsx', '.md'),
            ('.docx', '.md'),
            ('.pdf', '.md'),
            ('.xlsx', '.html'),
            ('.docx', '.html'),
            ('.xlsx', '.txt'),
            ('.docx', '.txt'),
            ('.pdf', '.txt')
        ]
        
        print("\n支持的转换:")
        for src, dst in conversions_to_test:
            supported = converter.can_convert(src, dst)
            status = "✓" if supported else "✗"
            print(f"  {status} {src} -> {dst}")
        
        return converter
    except Exception as e:
        print(f"✗ 转换器初始化失败: {e}")
        traceback.print_exc()
        return None

def create_test_files():
    """创建测试文件"""
    print("\n=== 创建测试文件 ===")
    test_files = {}
    
    # 创建Excel测试文件
    try:
        import pandas as pd
        data = {
            '姓名': ['张三', '李四'],
            '年龄': [25, 30],
            '部门': ['技术部', '销售部']
        }
        df = pd.DataFrame(data)
        excel_file = 'diagnostic_test.xlsx'
        df.to_excel(excel_file, index=False)
        test_files['excel'] = excel_file
        print(f"✓ 创建Excel测试文件: {excel_file}")
    except Exception as e:
        print(f"✗ 创建Excel文件失败: {e}")
    
    # 创建Word测试文件
    try:
        from docx import Document
        doc = Document()
        doc.add_heading('测试文档', 0)
        doc.add_paragraph('这是一个测试段落。')
        word_file = 'diagnostic_test.docx'
        doc.save(word_file)
        test_files['word'] = word_file
        print(f"✓ 创建Word测试文件: {word_file}")
    except Exception as e:
        print(f"✗ 创建Word文件失败: {e}")
    
    # 创建PDF测试文件
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font('Arial', 'B', 16)
        pdf.cell(40, 10, 'Test Document')
        pdf.ln(20)
        pdf.set_font('Arial', '', 12)
        pdf.cell(40, 10, 'This is a test paragraph.')
        pdf_file = 'diagnostic_test.pdf'
        pdf.output(pdf_file)
        test_files['pdf'] = pdf_file
        print(f"✓ 创建PDF测试文件: {pdf_file}")
    except Exception as e:
        print(f"✗ 创建PDF文件失败: {e}")
    
    return test_files

def test_conversions(converter, test_files):
    """测试所有转换"""
    print("\n=== 测试转换功能 ===")
    
    results = {}
    
    for file_type, file_path in test_files.items():
        print(f"\n--- 测试 {file_type.upper()} 转换 ---")
        
        if not os.path.exists(file_path):
            print(f"✗ 文件不存在: {file_path}")
            continue
        
        # 测试转换为Markdown
        try:
            result = converter.convert_file(file_path, '.md')
            if result and os.path.exists(result):
                print(f"✓ {file_type} -> Markdown: {result}")
                results[f"{file_type}_md"] = True
                
                # 检查内容
                with open(result, 'r', encoding='utf-8') as f:
                    content = f.read()
                    if content.strip():
                        print(f"  内容长度: {len(content)} 字符")
                    else:
                        print("  ⚠️  内容为空")
            else:
                print(f"✗ {file_type} -> Markdown 失败")
                results[f"{file_type}_md"] = False
        except Exception as e:
            print(f"✗ {file_type} -> Markdown 出错: {e}")
            results[f"{file_type}_md"] = False
    
    return results

def generate_optimization_suggestions(results):
    """生成优化建议"""
    print("\n=== 优化建议 ===")
    
    failed_conversions = [k for k, v in results.items() if not v]
    
    if not failed_conversions:
        print("🎉 所有转换都成功！系统运行正常。")
        return
    
    print("发现以下问题:")
    for failed in failed_conversions:
        print(f"  ✗ {failed}")
    
    print("\n建议的优化措施:")
    
    if any('excel' in f for f in failed_conversions):
        print("📊 Excel转换问题:")
        print("  - 确保安装了 pandas 和 openpyxl")
        print("  - 检查Excel文件格式是否正确")
        print("  - 尝试重新安装: pip install pandas openpyxl")
    
    if any('word' in f for f in failed_conversions):
        print("📄 Word转换问题:")
        print("  - 确保安装了 python-docx")
        print("  - 检查Word文件是否损坏")
        print("  - 尝试重新安装: pip install python-docx")
    
    if any('pdf' in f for f in failed_conversions):
        print("📋 PDF转换问题:")
        print("  - 确保安装了 pypdf 和 pdfplumber")
        print("  - 检查PDF文件是否有文本内容")
        print("  - 尝试重新安装: pip install pypdf pdfplumber")
    
    print("\n🔧 通用优化建议:")
    print("  - 更新所有依赖包: pip install --upgrade pandas openpyxl python-docx pypdf pdfplumber")
    print("  - 检查Pandoc版本: pandoc --version")
    print("  - 重启应用程序")
    print("  - 检查文件权限")

def cleanup_test_files(test_files):
    """清理测试文件"""
    print("\n=== 清理测试文件 ===")
    
    all_files = list(test_files.values())
    
    # 添加可能的输出文件
    for file_path in test_files.values():
        base_name = os.path.splitext(file_path)[0]
        all_files.extend([
            f"{base_name}.md",
            f"{base_name}.html",
            f"{base_name}.txt"
        ])
    
    for file_path in all_files:
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                print(f"✓ 删除: {file_path}")
        except Exception as e:
            print(f"✗ 删除失败 {file_path}: {e}")

def main():
    """主函数"""
    print("DocuFlow 转换功能诊断和优化工具")
    print("=" * 50)
    
    # 检查依赖
    deps_ok = check_dependencies()
    pandoc_ok = check_pandoc()
    
    if not deps_ok:
        print("\n❌ 依赖检查失败，请先安装缺失的包")
        return
    
    if not pandoc_ok:
        print("\n❌ Pandoc检查失败，请先安装Pandoc")
        return
    
    # 测试转换器
    converter = test_converter_initialization()
    if not converter:
        print("\n❌ 转换器初始化失败")
        return
    
    # 创建测试文件
    test_files = create_test_files()
    if not test_files:
        print("\n❌ 无法创建测试文件")
        return
    
    # 测试转换
    results = test_conversions(converter, test_files)
    
    # 生成建议
    generate_optimization_suggestions(results)
    
    # 清理文件
    cleanup_test_files(test_files)
    
    print("\n" + "=" * 50)
    print("诊断完成！")

if __name__ == "__main__":
    main()