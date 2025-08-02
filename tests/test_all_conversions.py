import unittest
import os
import sys
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from converter.document_converter import DocumentConverter

class TestAllConversions(unittest.TestCase):
    def setUp(self):
        self.converter = DocumentConverter()
        self.output_dir = Path("tests/output")
        self.output_dir.mkdir(exist_ok=True)

        # Create dummy files
        self.test_files = {
            "pdf": self.output_dir / "test.pdf",
            "excel": self.output_dir / "test.xlsx",
            "word": self.output_dir / "test.docx",
        }
        self._create_dummy_files()

    def _create_dummy_files(self):
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        import openpyxl
        from docx import Document

        # PDF
        c = canvas.Canvas(str(self.test_files["pdf"]), pagesize=letter)
        c.drawString(100, 750, "This is a test PDF.")
        c.save()

        # Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Header1"
        ws["B1"] = "Header2"
        ws["A2"] = "Data1"
        ws["B2"] = "Data2"
        wb.save(self.test_files["excel"])

        # Word
        doc = Document()
        doc.add_heading('Test Document', 1)
        doc.add_paragraph('This is a test Word document.')
        doc.save(self.test_files["word"])

    def test_pdf_to_md(self):
        output_path = self.converter.convert_file(str(self.test_files["pdf"]), ".md", str(self.output_dir))
        self.assertTrue(Path(output_path).exists())
        with open(output_path, 'r') as f:
            content = f.read()
            self.assertIn("This is a test PDF", content)

    def test_excel_to_md(self):
        output_path = self.converter.convert_file(str(self.test_files["excel"]), ".md", str(self.output_dir))
        self.assertTrue(Path(output_path).exists())
        with open(output_path, 'r') as f:
            content = f.read()
            self.assertIn("Header1", content)
            self.assertIn("Data2", content)

    def test_word_to_md(self):
        output_path = self.converter.convert_file(str(self.test_files["word"]), ".md", str(self.output_dir))
        self.assertTrue(Path(output_path).exists())
        with open(output_path, 'r', encoding='utf-8') as f:
            content = f.read()
            # With level 1 heading, Pandoc should produce a markdown heading
            self.assertIn("# Test Document", content)

if __name__ == '__main__':
    unittest.main()