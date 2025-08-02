#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
测试转换问题诊断脚本
重现用户反馈的Word转HTML中文乱码和HTML转Markdown文件找不到的问题
"""

import sys
import os
from pathlib import Path
import tempfile
import subprocess

# 添加项目根目录到Python路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from converter.document_converter import DocumentConverter
from utils.file_utils import is_supported_file

def create_test_word_file():
    """创建包含中文内容的测试Word文件"""
    print("🔧 创建包含中文内容的测试Word文件...")
    
    try:
        from docx import Document
        
        # 创建Word文档
        doc = Document()
        
        # 添加标题
        doc.add_heading('测试文档标题', 0)
        
        # 添加段落
        doc.add_paragraph('这是一个包含中文内容的测试文档。')
        doc.add_paragraph('本文档用于测试Word转HTML时的中文编码问题。')
        
        # 添加二级标题
        doc.add_heading('功能测试', level=1)
        
        # 添加列表
        doc.add_paragraph('测试功能包括：', style='List Bullet')
        doc.add_paragraph('• 中文字符显示', style='List Bullet')
        doc.add_paragraph('• 特殊符号：！@#￥%……&*（）', style='List Bullet')
        doc.add_paragraph('• 数字和英文：123 ABC test', style='List Bullet')
        
        # 添加表格
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = '姓名'
        hdr_cells[1].text = '年龄'
        hdr_cells[2].text = '城市'
        
        row_cells = table.add_row().cells
        row_cells[0].text = '张三'
        row_cells[1].text = '25'
        row_cells[2].text = '北京'
        
        row_cells = table.add_row().cells
        row_cells[0].text = '李四'
        row_cells[1].text = '30'
        row_cells[2].text = '上海'
        
        # 保存文件
        test_file = 'test_chinese_content.docx'
        doc.save(test_file)
        
        print(f"✅ 创建测试Word文件: {test_file}")
        return test_file
        
    except ImportError:
        print("⚠️ python-docx库未安装，无法创建Word文件")
        return None
    except Exception as e:
        print(f"❌ 创建Word文件时出错: {e}")
        return None

def create_test_html_file():
    """创建包含中文内容的测试HTML文件"""
    print("\n🔧 创建包含中文内容的测试HTML文件...")
    
    html_content = '''<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>测试HTML文档</title>
</head>
<body>
    <h1>测试HTML文档标题</h1>
    
    <p>这是一个包含中文内容的测试HTML文档。</p>
    <p>本文档用于测试HTML转Markdown时的文件输出问题。</p>
    
    <h2>功能测试</h2>
    
    <p>测试功能包括：</p>
    <ul>
        <li>中文字符显示</li>
        <li>特殊符号：！@#￥%……&*（）</li>
        <li>数字和英文：123 ABC test</li>
    </ul>
    
    <h3>表格测试</h3>
    <table border="1">
        <thead>
            <tr>
                <th>姓名</th>
                <th>年龄</th>
                <th>城市</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td>张三</td>
                <td>25</td>
                <td>北京</td>
            </tr>
            <tr>
                <td>李四</td>
                <td>30</td>
                <td>上海</td>
            </tr>
        </tbody>
    </table>
    
    <h3>代码块测试</h3>
    <pre><code>def hello():
    print("你好，世界！")
    return "Hello, World!"</code></pre>
    
    <p><strong>粗体文本</strong> 和 <em>斜体文本</em></p>
    
    <p>链接测试：<a href="https://example.com">示例链接</a></p>
</body>
</html>'''
    
    try:
        test_file = 'test_chinese_content.html'
        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        print(f"✅ 创建测试HTML文件: {test_file}")
        return test_file
        
    except Exception as e:
        print(f"❌ 创建HTML文件时出错: {e}")
        return None

def test_word_to_html_encoding():
    """测试Word转HTML的中文编码问题"""
    print("\n🧪 测试Word转HTML中文编码问题")
    print("=" * 50)
    
    # 创建测试Word文件
    word_file = create_test_word_file()
    if not word_file:
        print("❌ 无法创建测试Word文件，跳过测试")
        return False
    
    try:
        # 初始化转换器
        converter = DocumentConverter()
        
        # 转换Word到HTML
        print("\n📋 执行Word转HTML转换...")
        output_file = converter.convert_file(word_file, '.html')
        
        print(f"✅ 转换完成，输出文件: {output_file}")
        
        # 检查输出文件是否存在
        if not Path(output_file).exists():
            print("❌ 输出文件不存在")
            return False
        
        # 读取HTML文件内容并检查编码
        print("\n📋 检查HTML文件内容和编码...")
        
        # 尝试不同编码读取
        encodings = ['utf-8', 'gbk', 'gb2312', 'utf-16']
        content = None
        used_encoding = None
        
        for encoding in encodings:
            try:
                with open(output_file, 'r', encoding=encoding) as f:
                    content = f.read()
                used_encoding = encoding
                print(f"✅ 成功使用 {encoding} 编码读取文件")
                break
            except UnicodeDecodeError:
                print(f"⚠️ {encoding} 编码读取失败")
                continue
        
        if content is None:
            print("❌ 无法使用任何编码读取HTML文件")
            return False
        
        # 检查中文内容
        chinese_keywords = ['测试文档标题', '中文内容', '张三', '北京']
        missing_keywords = []
        
        for keyword in chinese_keywords:
            if keyword not in content:
                missing_keywords.append(keyword)
        
        if missing_keywords:
            print(f"❌ HTML文件中缺少中文关键词: {', '.join(missing_keywords)}")
            print(f"📄 文件内容预览 (前500字符):")
            print(content[:500])
            return False
        else:
            print("✅ HTML文件包含所有中文关键词")
            
        # 检查HTML编码声明
        if 'charset=utf-8' in content.lower() or 'charset="utf-8"' in content.lower():
            print("✅ HTML文件包含UTF-8编码声明")
        else:
            print("⚠️ HTML文件缺少UTF-8编码声明")
        
        return True
        
    except Exception as e:
        print(f"❌ Word转HTML测试失败: {e}")
        return False
    finally:
        # 清理测试文件
        cleanup_files([word_file, word_file.replace('.docx', '.html')])

def test_html_to_markdown_output():
    """测试HTML转Markdown的文件输出问题"""
    print("\n🧪 测试HTML转Markdown文件输出问题")
    print("=" * 50)
    
    # 创建测试HTML文件
    html_file = create_test_html_file()
    if not html_file:
        print("❌ 无法创建测试HTML文件，跳过测试")
        return False
    
    try:
        # 初始化转换器
        converter = DocumentConverter()
        
        # 测试不同的输出路径场景
        test_scenarios = [
            {
                'name': '默认输出目录',
                'output_dir': None,
                'keep_original_name': True
            },
            {
                'name': '指定输出目录',
                'output_dir': './output_test',
                'keep_original_name': True
            },
            {
                'name': '临时目录输出',
                'output_dir': tempfile.gettempdir(),
                'keep_original_name': False
            }
        ]
        
        results = []
        
        for scenario in test_scenarios:
            print(f"\n📋 测试场景: {scenario['name']}")
            
            try:
                # 确保输出目录存在
                if scenario['output_dir']:
                    Path(scenario['output_dir']).mkdir(parents=True, exist_ok=True)
                
                # 执行转换
                output_file = converter.convert_file(
                    html_file, 
                    '.md',
                    output_dir=scenario['output_dir'],
                    keep_original_name=scenario['keep_original_name']
                )
                
                print(f"✅ 转换完成，输出文件: {output_file}")
                
                # 检查文件是否存在
                if Path(output_file).exists():
                    file_size = Path(output_file).stat().st_size
                    print(f"✅ 输出文件存在，大小: {file_size} 字节")
                    
                    # 读取并检查内容
                    with open(output_file, 'r', encoding='utf-8') as f:
                        content = f.read()
                    
                    # 检查Markdown内容
                    if '# 测试HTML文档标题' in content or '测试HTML文档标题' in content:
                        print("✅ Markdown内容包含预期的标题")
                        results.append(True)
                    else:
                        print("⚠️ Markdown内容可能不完整")
                        print(f"📄 内容预览: {content[:200]}...")
                        results.append(False)
                    
                    # 清理输出文件
                    Path(output_file).unlink()
                    
                else:
                    print(f"❌ 输出文件不存在: {output_file}")
                    results.append(False)
                    
            except Exception as e:
                print(f"❌ 场景测试失败: {e}")
                results.append(False)
        
        # 汇总结果
        success_count = sum(results)
        total_count = len(results)
        
        print(f"\n📊 HTML转Markdown测试结果: {success_count}/{total_count} 成功")
        
        return success_count == total_count
        
    except Exception as e:
        print(f"❌ HTML转Markdown测试失败: {e}")
        return False
    finally:
        # 清理测试文件
        cleanup_files([html_file])
        # 清理测试输出目录
        output_test_dir = Path('./output_test')
        if output_test_dir.exists():
            import shutil
            shutil.rmtree(output_test_dir, ignore_errors=True)

def test_pandoc_encoding_support():
    """测试Pandoc的编码支持"""
    print("\n🧪 测试Pandoc编码支持")
    print("=" * 50)
    
    try:
        # 检查Pandoc版本
        result = subprocess.run(['pandoc', '--version'], 
                              capture_output=True, text=True, timeout=5)
        if result.returncode == 0:
            version_info = result.stdout.split('\n')[0]
            print(f"✅ Pandoc版本: {version_info}")
        else:
            print("❌ 无法获取Pandoc版本信息")
            return False
        
        # 测试Pandoc对中文的支持
        test_html = '''<!DOCTYPE html>
<html>
<head><meta charset="UTF-8"><title>测试</title></head>
<body><h1>中文测试</h1><p>这是中文内容测试。</p></body>
</html>'''
        
        # 创建临时HTML文件
        with tempfile.NamedTemporaryFile(mode='w', suffix='.html', 
                                       encoding='utf-8', delete=False) as f:
            f.write(test_html)
            temp_html = f.name
        
        # 创建临时Markdown文件路径
        temp_md = temp_html.replace('.html', '.md')
        
        try:
            # 执行Pandoc转换
            cmd = ['pandoc', temp_html, '-o', temp_md, '--wrap=none']
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=10)
            
            if result.returncode == 0:
                print("✅ Pandoc转换执行成功")
                
                # 检查输出文件
                if Path(temp_md).exists():
                    with open(temp_md, 'r', encoding='utf-8') as f:
                        md_content = f.read()
                    
                    if '中文测试' in md_content and '中文内容测试' in md_content:
                        print("✅ Pandoc正确处理中文内容")
                        return True
                    else:
                        print(f"❌ Pandoc中文处理异常，输出内容: {md_content}")
                        return False
                else:
                    print("❌ Pandoc未生成输出文件")
                    return False
            else:
                print(f"❌ Pandoc转换失败: {result.stderr}")
                return False
                
        finally:
            # 清理临时文件
            cleanup_files([temp_html, temp_md])
        
    except Exception as e:
        print(f"❌ Pandoc编码测试失败: {e}")
        return False

def cleanup_files(file_paths):
    """清理测试文件"""
    for file_path in file_paths:
        try:
            if file_path and Path(file_path).exists():
                Path(file_path).unlink()
        except Exception:
            pass  # 忽略清理错误

def main():
    print("🧪 转换问题诊断测试")
    print("=" * 60)
    
    test_results = []
    
    # 测试1: Pandoc编码支持
    result1 = test_pandoc_encoding_support()
    test_results.append(('Pandoc编码支持', result1))
    
    # 测试2: Word转HTML中文编码
    result2 = test_word_to_html_encoding()
    test_results.append(('Word转HTML中文编码', result2))
    
    # 测试3: HTML转Markdown文件输出
    result3 = test_html_to_markdown_output()
    test_results.append(('HTML转Markdown文件输出', result3))
    
    # 汇总结果
    print("\n" + "=" * 60)
    print("📊 诊断测试结果汇总:")
    
    passed_tests = 0
    total_tests = len(test_results)
    
    for test_name, result in test_results:
        status = "✅ 通过" if result else "❌ 失败"
        print(f"  {test_name}: {status}")
        if result:
            passed_tests += 1
    
    print(f"\n🎯 总体结果: {passed_tests}/{total_tests} 测试通过")
    
    if passed_tests < total_tests:
        print("\n⚠️ 发现问题，需要进行修复")
        print("\n🔧 建议的修复措施:")
        
        if not test_results[0][1]:  # Pandoc编码支持失败
            print("  1. 检查Pandoc安装和版本")
            print("  2. 确保系统支持UTF-8编码")
        
        if not test_results[1][1]:  # Word转HTML编码失败
            print("  3. 在Pandoc命令中添加编码参数")
            print("  4. 确保Word文件使用UTF-8编码")
        
        if not test_results[2][1]:  # HTML转Markdown输出失败
            print("  5. 检查输出路径权限")
            print("  6. 验证文件路径生成逻辑")
    else:
        print("\n🎉 所有测试通过！转换功能正常")
    
    return passed_tests == total_tests

if __name__ == '__main__':
    success = main()
    sys.exit(0 if success else 1)