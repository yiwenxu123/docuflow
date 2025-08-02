#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
创建演示用的 Excel 文件
"""

import sys
from pathlib import Path

# 尝试导入 openpyxl
try:
    import openpyxl
except ImportError:
    print("正在安装 openpyxl...")
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
    import openpyxl

def create_demo_excel():
    """创建演示用的 Excel 文件"""
    # 创建工作簿
    workbook = openpyxl.Workbook()
    
    # 删除默认工作表
    workbook.remove(workbook.active)
    
    # 创建员工信息工作表
    ws1 = workbook.create_sheet("员工信息")
    
    # 添加员工数据
    employee_data = [
        ["姓名", "部门", "职位", "工资", "入职日期"],
        ["张三", "技术部", "高级工程师", 15000, "2022-01-15"],
        ["李四", "产品部", "产品经理", 18000, "2021-08-20"],
        ["王五", "设计部", "UI设计师", 12000, "2023-03-10"],
        ["赵六", "技术部", "前端工程师", 13000, "2022-11-05"],
        ["钱七", "运营部", "运营专员", 10000, "2023-06-01"]
    ]
    
    for row_idx, row_data in enumerate(employee_data, 1):
        for col_idx, value in enumerate(row_data, 1):
            ws1.cell(row=row_idx, column=col_idx, value=value)
    
    # 创建项目统计工作表
    ws2 = workbook.create_sheet("项目统计")
    
    # 添加项目数据
    project_data = [
        ["项目名称", "负责人", "状态", "进度", "预算(万元)"],
        ["DocuFlow", "张三", "进行中", "85%", 50],
        ["移动应用", "李四", "已完成", "100%", 30],
        ["网站重构", "王五", "计划中", "10%", 20],
        ["数据分析", "赵六", "进行中", "60%", 15]
    ]
    
    for row_idx, row_data in enumerate(project_data, 1):
        for col_idx, value in enumerate(row_data, 1):
            ws2.cell(row=row_idx, column=col_idx, value=value)
    
    # 创建销售数据工作表
    ws3 = workbook.create_sheet("销售数据")
    
    # 添加销售数据
    sales_data = [
        ["月份", "产品A", "产品B", "产品C", "总计"],
        ["1月", 1200, 800, 600, 2600],
        ["2月", 1500, 900, 700, 3100],
        ["3月", 1800, 1200, 800, 3800],
        ["4月", 2000, 1100, 900, 4000]
    ]
    
    for row_idx, row_data in enumerate(sales_data, 1):
        for col_idx, value in enumerate(row_data, 1):
            ws3.cell(row=row_idx, column=col_idx, value=value)
    
    # 保存文件
    file_path = Path(__file__).parent / "demo_data.xlsx"
    workbook.save(file_path)
    workbook.close()
    
    print(f"✅ 演示 Excel 文件已创建: {file_path}")
    print(f"📊 包含 3 个工作表:")
    print(f"   - 员工信息 (6行数据)")
    print(f"   - 项目统计 (5行数据)")
    print(f"   - 销售数据 (5行数据)")
    
    return str(file_path)

if __name__ == "__main__":
    create_demo_excel()