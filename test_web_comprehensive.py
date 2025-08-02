#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Web版本综合功能测试脚本
测试多种文件格式的上传和转换功能
"""

import os
import requests
import json
import time
from docx import Document
from pathlib import Path

def create_test_files():
    """创建各种格式的测试文件"""
    print("📁 创建测试文件...")
    
    # 创建Word文档
    doc = Document()
    doc.add_heading('测试文档', 0)
    doc.add_paragraph('这是一个测试段落，用于验证文档转换功能。')
    doc.add_paragraph('包含中文内容以测试编码处理。')
    doc.save('test_word.docx')
    print("✅ 创建: test_word.docx")
    
    # 创建Markdown文件
    md_content = """# 测试Markdown文档

## 简介
这是一个测试Markdown文档，用于验证转换功能。

## 内容
- 列表项1
- 列表项2
- 列表项3

### 代码示例
```python
def hello():
    print("Hello, World!")
```

**粗体文本** 和 *斜体文本*
"""
    with open('test_markdown.md', 'w', encoding='utf-8') as f:
        f.write(md_content)
    print("✅ 创建: test_markdown.md")
    
    # 创建HTML文件
    html_content = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>测试HTML文档</title>
</head>
<body>
    <h1>测试HTML文档</h1>
    <p>这是一个测试HTML文档，用于验证转换功能。</p>
    <ul>
        <li>列表项1</li>
        <li>列表项2</li>
        <li>列表项3</li>
    </ul>
    <p><strong>粗体文本</strong> 和 <em>斜体文本</em></p>
</body>
</html>"""
    with open('test_html.html', 'w', encoding='utf-8') as f:
        f.write(html_content)
    print("✅ 创建: test_html.html")

def test_server_connection():
    """测试服务器连接"""
    try:
        response = requests.get('http://localhost:5003/', timeout=5)
        if response.status_code == 200:
            print("✅ 服务器连接正常")
            return True
        else:
            print(f"❌ 服务器响应异常: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 服务器连接失败: {e}")
        return False

def test_file_upload_and_conversion(file_path, target_format):
    """测试文件上传和转换"""
    print(f"\n🔄 测试: {file_path} -> {target_format}")
    
    # 上传文件
    try:
        with open(file_path, 'rb') as f:
            files = {'files': (os.path.basename(file_path), f)}
            upload_response = requests.post('http://localhost:5003/upload', files=files, timeout=10)
        
        if upload_response.status_code != 200:
            print(f"❌ 上传失败: {upload_response.status_code}")
            return False
            
        upload_data = upload_response.json()
        if not upload_data.get('success'):
            print(f"❌ 上传失败: {upload_data.get('message', '未知错误')}")
            return False
            
        print(f"✅ 上传成功: {upload_data['message']}")
        
        # 转换文件
        file_paths = [file_info['path'] for file_info in upload_data['files']]
        convert_data = {
            'file_paths': file_paths,
            'output_format': target_format,
            'keep_original_name': True,
            'use_temp_dir': True
        }
        
        convert_response = requests.post('http://localhost:5003/convert', 
                                       json=convert_data, timeout=30)
        
        if convert_response.status_code != 200:
            print(f"❌ 转换失败: {convert_response.status_code}")
            return False
            
        convert_result = convert_response.json()
        if not convert_result.get('success'):
            print(f"❌ 转换失败: {convert_result}")
            return False
            
        success_count = convert_result['summary']['success']
        total_count = convert_result['summary']['total']
        print(f"✅ 转换成功: {success_count}/{total_count}")
        
        return True
        
    except Exception as e:
        print(f"❌ 测试失败: {e}")
        return False

def test_statistics():
    """测试统计功能"""
    try:
        response = requests.get('http://localhost:5003/stats', timeout=5)
        if response.status_code == 200:
            stats = response.json()
            print(f"✅ 统计功能正常: 总转换数 {stats['stats']['total_conversions']}")
            return True
        else:
            print(f"❌ 统计功能异常: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 统计功能失败: {e}")
        return False

def cleanup_test_files():
    """清理测试文件"""
    test_files = ['test_word.docx', 'test_markdown.md', 'test_html.html']
    for file in test_files:
        if os.path.exists(file):
            os.remove(file)
            print(f"🗑️ 删除: {file}")

def main():
    """主测试函数"""
    print("🧪 开始Web版本综合功能测试...")
    print("=" * 50)
    
    # 创建测试文件
    create_test_files()
    
    # 测试服务器连接
    if not test_server_connection():
        print("❌ 服务器连接失败，测试终止")
        cleanup_test_files()
        return
    
    # 测试用例
    test_cases = [
        ('test_word.docx', '.md'),
        ('test_word.docx', '.html'),
        ('test_markdown.md', '.docx'),
        ('test_markdown.md', '.html'),
        ('test_html.html', '.docx'),
        ('test_html.html', '.md'),
    ]
    
    success_count = 0
    total_count = len(test_cases)
    
    # 执行测试
    for file_path, target_format in test_cases:
        if test_file_upload_and_conversion(file_path, target_format):
            success_count += 1
        time.sleep(1)  # 避免请求过快
    
    # 测试统计功能
    stats_ok = test_statistics()
    
    # 输出测试结果
    print("\n" + "=" * 50)
    print("📋 测试总结")
    print("=" * 50)
    print(f"🎯 转换测试: {success_count}/{total_count} 成功")
    print(f"📊 统计功能: {'✅ 正常' if stats_ok else '❌ 异常'}")
    
    if success_count == total_count and stats_ok:
        print("🎉 总体结果: ✅ 所有功能正常")
    else:
        print("⚠️ 总体结果: ❌ 部分功能异常")
    
    # 清理测试文件
    print("\n🧹 清理测试文件...")
    cleanup_test_files()

if __name__ == "__main__":
    main()