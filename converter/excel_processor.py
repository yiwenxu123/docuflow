#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 处理器 Mixin

负责 Excel 数据提取与 Excel 转 Markdown/HTML/TXT 的专用转换逻辑。
"""

import os
from pathlib import Path
from typing import List, Any, Union

try:
    import openpyxl
    import pandas as pd

    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False

from exceptions import ConversionError


class ExcelProcessorMixin:
    """Excel 提取与转换逻辑"""

    def _extract_data_from_excel(self, excel_path: Union[str, Path]) -> str:
        """从Excel文件中提取数据并转换为Markdown，增加详细日志和错误处理。

        Args:
            excel_path: Excel文件路径

        Returns:
            str: Markdown格式的内容

        Raises:
            ConversionError: 当Excel处理库未安装或数据提取失败时
        """
        if not EXCEL_SUPPORT:
            raise ConversionError(
                "Excel处理库 (openpyxl, pandas) 未安装", error_code="EXCEL_LIBS_MISSING"
            )

        extraction_methods = [
            ("openpyxl", self._extract_with_openpyxl),
            ("pandas", self._extract_with_pandas),
        ]

        for name, method in extraction_methods:
            try:
                self.logger.info(f"尝试使用 {name} 提取Excel数据: {excel_path}")
                content = method(excel_path)
                if content and content.strip():
                    self.logger.info(f"成功使用 {name} 提取到内容")
                    return content
                self.logger.warning(f"使用 {name} 未提取到有效内容")
            except Exception as e:
                self.logger.error(f"使用 {name} 提取Excel失败: {e}", exc_info=True)

        raise ConversionError(
            "所有Excel数据提取方法均失败", error_code="EXCEL_EXTRACTION_FAILED"
        )

    def _extract_with_openpyxl(self, excel_path: Union[str, Path]) -> str:
        """使用openpyxl从Excel中提取数据。"""
        workbook = openpyxl.load_workbook(excel_path, data_only=True)
        markdown_content = [
            f"# {os.path.splitext(os.path.basename(excel_path))[0]}",
            "\n> 本文档由DocuFlow从Excel自动转换生成\n",
        ]

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            if len(workbook.sheetnames) > 1:
                markdown_content.append(f"## 工作表: {sheet_name}\n")

            data = []
            for row in worksheet.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() for cell in row):
                    cleaned_row = [
                        str(cell).strip() if cell is not None else "" for cell in row
                    ]
                    data.append(cleaned_row)

            if not data:
                markdown_content.append("*此工作表没有有效数据*\n")
                continue

            markdown_content.append(self._convert_table_to_markdown(data))
            markdown_content.append(
                f"\n*数据行数: {len(data)}, 列数: {len(data[0]) if data else 0}*\n"
            )

        workbook.close()
        return "\n".join(markdown_content)

    def _extract_with_pandas(self, excel_path: Union[str, Path]) -> str:
        """使用pandas从Excel中提取数据作为备用方法。"""
        excel_data = pd.read_excel(excel_path, sheet_name=None, engine="openpyxl")
        markdown_content = [
            f"# {os.path.splitext(os.path.basename(excel_path))[0]}",
            "\n> 本文档由DocuFlow从Excel自动转换生成\n",
        ]

        for sheet_name, df in excel_data.items():
            if len(excel_data) > 1:
                markdown_content.append(f"## 工作表: {sheet_name}\n")

            if df.empty:
                markdown_content.append("*此工作表为空*\n")
                continue

            df_cleaned = df.dropna(how="all").dropna(axis=1, how="all").fillna("")
            if df_cleaned.empty:
                markdown_content.append("*此工作表没有有效数据*\n")
                continue

            data = [df_cleaned.columns.astype(str).tolist()] + df_cleaned.astype(
                str
            ).values.tolist()
            markdown_content.append(self._convert_table_to_markdown(data))
            markdown_content.append(
                f"\n*数据行数: {len(df_cleaned)}, 列数: {len(df_cleaned.columns)}*\n"
            )

        return "\n".join(markdown_content)

    def _convert_excel_to_markdown(
        self, excel_path: Union[str, Path], output_path: Union[str, Path]
    ) -> str:
        """将Excel转换为Markdown

        Args:
            excel_path: Excel文件路径
            output_path: 输出文件路径

        Returns:
            str: 输出文件路径
        """
        # 提取Excel数据内容
        markdown_content = self._extract_data_from_excel(excel_path)

        # 写入输出文件
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(markdown_content)

        return output_path

    def _convert_excel_to_html(
        self, excel_path: Union[str, Path], output_path: Union[str, Path]
    ) -> str:
        """将Excel转换为HTML

        Args:
            excel_path: Excel文件路径
            output_path: 输出文件路径

        Returns:
            str: 输出文件路径
        """
        # 提取Excel数据内容（Markdown格式）
        markdown_content = self._extract_data_from_excel(excel_path)

        # 将Markdown转换为HTML
        html_content = self._convert_markdown_to_html(markdown_content)

        # 写入输出文件
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html_content)

        return output_path

    def _convert_excel_to_txt(
        self, excel_path: Union[str, Path], output_path: Union[str, Path]
    ) -> str:
        """将Excel转换为TXT

        Args:
            excel_path: Excel文件路径
            output_path: 输出文件路径

        Returns:
            str: 输出文件路径
        """
        # 提取Excel数据内容（Markdown格式）
        markdown_content = self._extract_data_from_excel(excel_path)

        # 将Markdown转换为纯文本
        txt_content = self._convert_markdown_to_txt(markdown_content)

        # 写入输出文件
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(txt_content)

        return output_path
