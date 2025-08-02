#!/usr/bin/env python3
"""
全面诊断工具 - 检查文档转换系统的所有组件
"""

import os
import sys
import subprocess
import traceback
from pathlib import Path
import tempfile
import json

def print_section(title):
    """打印分节标题"""
    print(f"\n{'='*60}")
    print(f"🔍 {title}")
    print('='*60)

def check_python_environment():
    """检查Python环境"""
    print_section("Python环境检查")
    
    print(f"Python版本: {sys.version}")
    print(f"Python路径: {sys.executable}")
    print(f"当前工作目录: {os.getcwd()}")
    print(f"Python路径: {sys.path[:3]}...")

def check_dependencies():
    """检查依赖包"""
    print_section("依赖包检查")
    
    required_packages = [
        ('pandas', 'pandas'), 
        ('openpyxl', 'openpyxl'), 
        ('python-docx', 'docx'), 
        ('PyPDF2', 'PyPDF2'),
        ('pdfplumber', 'pdfplumber'), 
        ('flask', 'flask'), 
        ('PySide6', 'PySide6'), 
        ('markdown', 'markdown')
    ]
    
    missing_packages = []
    
    for package_name, import_name in required_packages:
        try:
            __import__(import_name)
            print(f"✅ {package_name}: 已安装")
        except ImportError:
            print(f"❌ {package_name}: 未安装")
            missing_packages.append(package_name)
    
    if missing_packages:
        print(f"\n⚠️  缺少依赖包: {', '.join(missing_packages)}")
        return False
    return True

def check_external_tools():
    """检查外部工具"""
    print_section("外部工具检查")
    
    tools = ['pandoc']
    all_available = True
    
    for tool in tools:
        try:
            result = subprocess.run(['which', tool], capture_output=True, text=True)
            if result.returncode == 0:
                path = result.stdout.strip()
                print(f"✅ {tool}: {path}")
                
                # 检查版本
                try:
                    version_result = subprocess.run([tool, '--version'], capture_output=True, text=True)
                    if version_result.returncode == 0:
                        version = version_result.stdout.split('\n')[0]
                        print(f"   版本: {version}")
                except:
                    print(f"   无法获取版本信息")
            else:
                print(f"❌ {tool}: 未找到")
                all_available = False
        except Exception as e:
            print(f"❌ {tool}: 检查失败 - {e}")
            all_available = False
    
    return all_available

def check_file_permissions():
    """检查文件权限"""
    print_section("文件权限检查")
    
    current_dir = Path.cwd()
    temp_dir = Path(tempfile.gettempdir())
    
    # 检查当前目录权限
    try:
        test_file = current_dir / "test_write_permission.tmp"
        test_file.write_text("test")
        test_file.unlink()
        print(f"✅ 当前目录写权限: {current_dir}")
    except Exception as e:
        print(f"❌ 当前目录写权限失败: {e}")
        return False
    
    # 检查临时目录权限
    try:
        test_file = temp_dir / "test_write_permission.tmp"
        test_file.write_text("test")
        test_file.unlink()
        print(f"✅ 临时目录写权限: {temp_dir}")
    except Exception as e:
        print(f"❌ 临时目录写权限失败: {e}")
        return False
    
    return True

def test_converter_import():
    """测试转换器导入"""
    print_section("转换器模块检查")
    
    try:
        # 添加当前目录到Python路径
        sys.path.insert(0, str(Path.cwd()))
        
        from converter.document_converter import DocumentConverter
        print("✅ DocumentConverter导入成功")
        
        # 尝试初始化
        converter = DocumentConverter()
        print("✅ DocumentConverter初始化成功")
        
        return converter
    except Exception as e:
        print(f"❌ DocumentConverter导入/初始化失败: {e}")
        traceback.print_exc()
        return None

def test_config_system():
    """测试配置系统"""
    print_section("配置系统检查")
    
    try:
        from utils.config_manager import get_config_manager
        
        config_manager = get_config_manager()
        print("✅ 配置管理器初始化成功")
        
        supported_conversions = config_manager.SUPPORTED_CONVERSIONS
        print(f"✅ 支持的转换格式: {len(supported_conversions)} 种")
        
        for source, targets in supported_conversions.items():
            print(f"   {source} -> {targets}")
        
        return config_manager
    except Exception as e:
        print(f"❌ 配置系统失败: {e}")
        traceback.print_exc()
        return None

def create_test_files():
    """创建测试文件"""
    print_section("创建测试文件")
    
    test_files = {}
    
    # 创建测试Excel文件
    try:
        import pandas as pd
        df = pd.DataFrame({
            'Name': ['张三', '李四', '王五'],
            'Age': [25, 30, 35],
            'City': ['北京', '上海', '广州']
        })
        excel_path = Path('test_diagnostic.xlsx')
        df.to_excel(excel_path, index=False)
        test_files['excel'] = excel_path
        print(f"✅ 创建测试Excel文件: {excel_path}")
    except Exception as e:
        print(f"❌ 创建Excel文件失败: {e}")
    
    # 创建测试Word文件
    try:
        from docx import Document
        doc = Document()
        doc.add_heading('测试文档', 0)
        doc.add_paragraph('这是一个测试段落。')
        doc.add_paragraph('包含中文内容的测试。')
        
        word_path = Path('test_diagnostic.docx')
        doc.save(word_path)
        test_files['word'] = word_path
        print(f"✅ 创建测试Word文件: {word_path}")
    except Exception as e:
        print(f"❌ 创建Word文件失败: {e}")
    
    # 创建测试Markdown文件
    try:
        md_content = """# 测试文档

这是一个测试Markdown文件。

## 功能列表

- 项目1
- 项目2
- 项目3

**粗体文本** 和 *斜体文本*
"""
        md_path = Path('test_diagnostic.md')
        md_path.write_text(md_content, encoding='utf-8')
        test_files['markdown'] = md_path
        print(f"✅ 创建测试Markdown文件: {md_path}")
    except Exception as e:
        print(f"❌ 创建Markdown文件失败: {e}")
    
    return test_files

def test_conversions(converter, test_files):
    """测试转换功能"""
    print_section("转换功能测试")
    
    if not converter or not test_files:
        print("❌ 无法进行转换测试 - 转换器或测试文件不可用")
        return
    
    test_cases = [
        ('excel', 'xlsx', 'md'),
        ('excel', 'xlsx', 'html'),
        ('word', 'docx', 'md'),
        ('word', 'docx', 'html'),
        ('markdown', 'md', 'html'),
        ('markdown', 'md', 'docx'),
    ]
    
    results = []
    output_files = []  # 记录所有输出文件，稍后清理
    
    for file_type, input_ext, output_ext in test_cases:
        if file_type not in test_files:
            print(f"⏭️  跳过 {file_type} -> {output_ext}: 测试文件不存在")
            continue
        
        input_file = test_files[file_type]
        output_file = input_file.with_suffix(f'.{output_ext}')
        
        try:
            print(f"🔄 测试 {input_file.name} -> {output_file.name}")
            
            # 执行转换
            result_path = converter.convert_file(
                str(input_file),
                str(output_file.resolve())  # 传递绝对路径作为output_format
            )
            
            if result_path and output_file.exists():
                file_size = output_file.stat().st_size
                print(f"✅ {file_type} -> {output_ext}: 成功 ({file_size} bytes)")
                results.append((file_type, output_ext, True, file_size))
                output_files.append(output_file)  # 记录输出文件
            else:
                print(f"❌ {file_type} -> {output_ext}: 失败 - 输出文件不存在或转换返回空")
                print(f"   result_path: {result_path}")
                print(f"   output_file.exists(): {output_file.exists()}")
                results.append((file_type, output_ext, False, 0))
                
        except Exception as e:
            print(f"❌ {file_type} -> {output_ext}: 异常 - {e}")
            print(f"   异常类型: {type(e).__name__}")
            import traceback
            traceback.print_exc()
            results.append((file_type, output_ext, False, 0))
    
    # 清理所有输出文件
    for output_file in output_files:
        try:
            if output_file.exists():
                output_file.unlink()
        except Exception as e:
            print(f"⚠️  清理输出文件失败 {output_file}: {e}")
    
    return results

def cleanup_test_files(test_files):
    """清理测试文件"""
    print_section("清理测试文件")
    
    for file_type, file_path in test_files.items():
        try:
            if file_path.exists():
                file_path.unlink()
                print(f"✅ 删除: {file_path}")
        except Exception as e:
            print(f"❌ 删除失败 {file_path}: {e}")

def generate_report(results):
    """生成诊断报告"""
    print_section("诊断报告")
    
    if not results:
        print("❌ 没有转换测试结果")
        return
    
    total_tests = len(results)
    successful_tests = sum(1 for _, _, success, _ in results if success)
    success_rate = (successful_tests / total_tests) * 100
    
    print(f"总测试数: {total_tests}")
    print(f"成功数: {successful_tests}")
    print(f"失败数: {total_tests - successful_tests}")
    print(f"成功率: {success_rate:.1f}%")
    
    print("\n详细结果:")
    for file_type, output_ext, success, size in results:
        status = "✅" if success else "❌"
        size_info = f"({size} bytes)" if success else ""
        print(f"  {status} {file_type} -> {output_ext} {size_info}")
    
    if success_rate < 100:
        print(f"\n⚠️  检测到转换问题，成功率仅为 {success_rate:.1f}%")
        print("建议检查:")
        print("1. 确保所有依赖包正确安装")
        print("2. 检查pandoc是否正常工作")
        print("3. 验证文件权限设置")
        print("4. 查看详细错误日志")
    else:
        print(f"\n🎉 所有转换测试通过！系统运行正常。")

def main():
    """主函数"""
    print("🚀 DocuFlow 全面诊断工具")
    print("正在检查系统状态...")
    
    # 基础环境检查
    check_python_environment()
    deps_ok = check_dependencies()
    tools_ok = check_external_tools()
    perms_ok = check_file_permissions()
    
    # 模块检查
    converter = test_converter_import()
    config_manager = test_config_system()
    
    # 如果基础检查都通过，进行转换测试
    if deps_ok and tools_ok and perms_ok and converter:
        test_files = create_test_files()
        results = test_conversions(converter, test_files)
        cleanup_test_files(test_files)
        generate_report(results)
    else:
        print_section("诊断结果")
        print("❌ 基础环境检查未通过，无法进行转换测试")
        print("请先解决上述问题后重新运行诊断")

if __name__ == "__main__":
    main()