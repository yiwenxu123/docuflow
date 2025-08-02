#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
全面项目检查和优化工具
检查所有功能模块的完整性和性能
"""

import os
import sys
import time
import json
import requests
import tempfile
import traceback
from pathlib import Path
from docx import Document
import pandas as pd

# 添加项目根目录到路径
project_root = Path(__file__).parent
sys.path.insert(0, str(project_root))

from converter.document_converter import DocumentConverter

class ProjectChecker:
    def __init__(self):
        self.results = {
            'core_functionality': {},
            'web_interface': {},
            'file_support': {},
            'performance': {},
            'error_handling': {},
            'overall_status': 'unknown'
        }
        self.converter = DocumentConverter()
        self.web_base_url = "http://localhost:5003"
        
    def log(self, message, level="INFO"):
        """日志记录"""
        timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
        print(f"[{timestamp}] [{level}] {message}")
        
    def check_dependencies(self):
        """检查依赖项"""
        self.log("检查项目依赖...")
        try:
            required_packages = [
                'flask', 'docx2txt', 'python-docx', 'markdownify', 
                'markdown', 'pandas', 'openpyxl', 'requests'
            ]
            
            missing_packages = []
            for package in required_packages:
                try:
                    __import__(package.replace('-', '_'))
                except ImportError:
                    missing_packages.append(package)
            
            if missing_packages:
                self.results['core_functionality']['dependencies'] = {
                    'status': 'warning',
                    'missing': missing_packages
                }
                self.log(f"缺少依赖包: {missing_packages}", "WARNING")
            else:
                self.results['core_functionality']['dependencies'] = {
                    'status': 'success',
                    'message': '所有依赖包已安装'
                }
                self.log("所有依赖包检查通过")
                
        except Exception as e:
            self.results['core_functionality']['dependencies'] = {
                'status': 'error',
                'error': str(e)
            }
            self.log(f"依赖检查失败: {e}", "ERROR")
    
    def check_core_conversion(self):
        """检查核心转换功能"""
        self.log("检查核心转换功能...")
        
        # 创建测试文件
        test_files = self.create_test_files()
        conversion_results = {}
        
        # 测试各种转换
        conversions = [
            ('docx', 'md'),
            ('docx', 'html'),
            ('md', 'html'),
            ('md', 'docx'),
            ('html', 'md'),
            ('html', 'docx'),
            ('xlsx', 'md')
        ]
        
        for source_format, target_format in conversions:
            try:
                if source_format in test_files:
                    source_file = test_files[source_format]
                    output_dir = tempfile.mkdtemp()
                    
                    start_time = time.time()
                    result = self.converter.convert_file(
                        source_file, f'.{target_format}', output_dir
                    )
                    conversion_time = time.time() - start_time
                    
                    if result and os.path.exists(result):
                        conversion_results[f'{source_format}_to_{target_format}'] = {
                            'status': 'success',
                            'time': round(conversion_time, 2),
                            'output_size': os.path.getsize(result)
                        }
                        self.log(f"✅ {source_format} → {target_format} 转换成功 ({conversion_time:.2f}s)")
                    else:
                        conversion_results[f'{source_format}_to_{target_format}'] = {
                            'status': 'failed',
                            'error': '输出文件未生成'
                        }
                        self.log(f"❌ {source_format} → {target_format} 转换失败", "ERROR")
                        
            except Exception as e:
                conversion_results[f'{source_format}_to_{target_format}'] = {
                    'status': 'error',
                    'error': str(e)
                }
                self.log(f"❌ {source_format} → {target_format} 转换异常: {e}", "ERROR")
        
        self.results['core_functionality']['conversions'] = conversion_results
        
        # 清理测试文件
        self.cleanup_test_files(test_files)
    
    def check_web_interface(self):
        """检查Web界面功能"""
        self.log("检查Web界面功能...")
        
        try:
            # 检查服务器连接
            response = requests.get(self.web_base_url, timeout=5)
            if response.status_code == 200:
                self.results['web_interface']['server_status'] = {
                    'status': 'success',
                    'message': 'Web服务器运行正常'
                }
                self.log("✅ Web服务器连接正常")
            else:
                self.results['web_interface']['server_status'] = {
                    'status': 'warning',
                    'code': response.status_code
                }
                self.log(f"⚠️ Web服务器响应异常: {response.status_code}", "WARNING")
                
        except requests.exceptions.ConnectionError:
            self.results['web_interface']['server_status'] = {
                'status': 'error',
                'error': '无法连接到Web服务器'
            }
            self.log("❌ 无法连接到Web服务器", "ERROR")
            return
        except Exception as e:
            self.results['web_interface']['server_status'] = {
                'status': 'error',
                'error': str(e)
            }
            self.log(f"❌ Web服务器检查异常: {e}", "ERROR")
            return
        
        # 测试API端点
        endpoints = ['/stats', '/cleanup']
        for endpoint in endpoints:
            try:
                if endpoint == '/cleanup':
                    response = requests.post(f"{self.web_base_url}{endpoint}", timeout=5)
                else:
                    response = requests.get(f"{self.web_base_url}{endpoint}", timeout=5)
                    
                if response.status_code == 200:
                    self.log(f"✅ API端点 {endpoint} 正常")
                else:
                    self.log(f"⚠️ API端点 {endpoint} 响应异常: {response.status_code}", "WARNING")
                    
            except Exception as e:
                self.log(f"❌ API端点 {endpoint} 检查失败: {e}", "ERROR")
        
        # 测试文件上传和转换
        self.test_web_upload_conversion()
    
    def test_web_upload_conversion(self):
        """测试Web文件上传和转换"""
        self.log("测试Web文件上传和转换...")
        
        try:
            # 创建测试文档
            test_doc_path = tempfile.mktemp(suffix='.docx')
            doc = Document()
            doc.add_heading('测试文档', 0)
            doc.add_paragraph('这是一个用于测试Web上传功能的文档。')
            doc.save(test_doc_path)
            
            # 上传文件
            with open(test_doc_path, 'rb') as f:
                files = {'files': f}
                upload_response = requests.post(
                    f"{self.web_base_url}/upload",
                    files=files,
                    timeout=10
                )
            
            if upload_response.status_code == 200:
                upload_data = upload_response.json()
                if upload_data.get('success'):
                    uploaded_files = upload_data.get('files', [])
                    self.log(f"✅ 文件上传成功: {len(uploaded_files)} 个文件")
                    
                    # 测试转换
                    if uploaded_files:
                        file_paths = [f['path'] for f in uploaded_files]
                        convert_data = {
                            'file_paths': file_paths,
                            'output_format': '.md'
                        }
                        
                        convert_response = requests.post(
                            f"{self.web_base_url}/convert",
                            json=convert_data,
                            timeout=15
                        )
                        
                        if convert_response.status_code == 200:
                            convert_result = convert_response.json()
                            if convert_result.get('success'):
                                self.results['web_interface']['upload_conversion'] = {
                                    'status': 'success',
                                    'message': '文件上传和转换功能正常'
                                }
                                self.log("✅ Web文件上传和转换功能正常")
                            else:
                                self.results['web_interface']['upload_conversion'] = {
                                    'status': 'error',
                                    'error': convert_result.get('error', '转换失败')
                                }
                                self.log(f"❌ Web转换失败: {convert_result.get('error')}", "ERROR")
                        else:
                            self.log(f"❌ 转换请求失败: {convert_response.status_code}", "ERROR")
                else:
                    self.log(f"❌ 文件上传失败: {upload_data.get('error')}", "ERROR")
            else:
                self.log(f"❌ 上传请求失败: {upload_response.status_code}", "ERROR")
                
            # 清理测试文件
            if os.path.exists(test_doc_path):
                os.remove(test_doc_path)
                
        except Exception as e:
            self.results['web_interface']['upload_conversion'] = {
                'status': 'error',
                'error': str(e)
            }
            self.log(f"❌ Web上传转换测试异常: {e}", "ERROR")
    
    def check_file_support(self):
        """检查文件格式支持"""
        self.log("检查文件格式支持...")
        
        supported_formats = {
            'input': ['.docx', '.md', '.html', '.htm', '.pdf', '.xlsx', '.xls'],
            'output': ['.docx', '.md', '.html', '.pdf']
        }
        
        format_check = {}
        
        # 检查输入格式支持
        for fmt in supported_formats['input']:
            try:
                can_convert = self.converter.can_convert_format(fmt)
                format_check[f'input_{fmt}'] = {
                    'supported': can_convert,
                    'status': 'success' if can_convert else 'warning'
                }
                status = "✅" if can_convert else "⚠️"
                self.log(f"{status} 输入格式 {fmt}: {'支持' if can_convert else '不支持'}")
            except Exception as e:
                format_check[f'input_{fmt}'] = {
                    'supported': False,
                    'status': 'error',
                    'error': str(e)
                }
                self.log(f"❌ 检查输入格式 {fmt} 时出错: {e}", "ERROR")
        
        self.results['file_support'] = format_check
    
    def check_performance(self):
        """检查性能指标"""
        self.log("检查性能指标...")
        
        try:
            # 创建大文件测试
            large_doc_path = tempfile.mktemp(suffix='.docx')
            doc = Document()
            doc.add_heading('性能测试文档', 0)
            
            # 添加大量内容
            for i in range(100):
                doc.add_paragraph(f'这是第 {i+1} 段测试内容。' * 10)
            
            doc.save(large_doc_path)
            file_size = os.path.getsize(large_doc_path)
            
            # 测试转换性能
            output_dir = tempfile.mkdtemp()
            start_time = time.time()
            
            result = self.converter.convert_file(large_doc_path, '.md', output_dir)
            
            conversion_time = time.time() - start_time
            
            if result and os.path.exists(result):
                output_size = os.path.getsize(result)
                self.results['performance']['large_file_conversion'] = {
                    'status': 'success',
                    'input_size_kb': round(file_size / 1024, 2),
                    'output_size_kb': round(output_size / 1024, 2),
                    'conversion_time_seconds': round(conversion_time, 2),
                    'speed_kb_per_second': round((file_size / 1024) / conversion_time, 2)
                }
                self.log(f"✅ 大文件转换性能: {file_size/1024:.1f}KB → {output_size/1024:.1f}KB 用时 {conversion_time:.2f}s")
            else:
                self.results['performance']['large_file_conversion'] = {
                    'status': 'failed',
                    'error': '大文件转换失败'
                }
                self.log("❌ 大文件转换失败", "ERROR")
            
            # 清理测试文件
            if os.path.exists(large_doc_path):
                os.remove(large_doc_path)
                
        except Exception as e:
            self.results['performance']['large_file_conversion'] = {
                'status': 'error',
                'error': str(e)
            }
            self.log(f"❌ 性能测试异常: {e}", "ERROR")
    
    def check_error_handling(self):
        """检查错误处理"""
        self.log("检查错误处理...")
        
        error_tests = []
        
        # 测试无效文件
        try:
            invalid_file = tempfile.mktemp(suffix='.docx')
            with open(invalid_file, 'w') as f:
                f.write("这不是一个有效的docx文件")
            
            output_dir = tempfile.mkdtemp()
            result = self.converter.convert_file(invalid_file, '.md', output_dir)
            
            if result is None:
                error_tests.append({
                    'test': 'invalid_file_handling',
                    'status': 'success',
                    'message': '正确处理了无效文件'
                })
                self.log("✅ 无效文件错误处理正常")
            else:
                error_tests.append({
                    'test': 'invalid_file_handling',
                    'status': 'warning',
                    'message': '无效文件未被正确识别'
                })
                self.log("⚠️ 无效文件处理可能有问题", "WARNING")
            
            os.remove(invalid_file)
            
        except Exception as e:
            error_tests.append({
                'test': 'invalid_file_handling',
                'status': 'success',
                'message': f'正确抛出异常: {str(e)[:50]}'
            })
            self.log("✅ 无效文件正确抛出异常")
        
        # 测试不存在的文件
        try:
            nonexistent_file = "/path/to/nonexistent/file.docx"
            output_dir = tempfile.mkdtemp()
            result = self.converter.convert_file(nonexistent_file, '.md', output_dir)
            
            error_tests.append({
                'test': 'nonexistent_file_handling',
                'status': 'warning',
                'message': '不存在的文件未正确处理'
            })
            self.log("⚠️ 不存在文件的处理可能有问题", "WARNING")
            
        except Exception as e:
            error_tests.append({
                'test': 'nonexistent_file_handling',
                'status': 'success',
                'message': f'正确处理不存在的文件: {str(e)[:50]}'
            })
            self.log("✅ 不存在文件错误处理正常")
        
        self.results['error_handling']['tests'] = error_tests
    
    def create_test_files(self):
        """创建测试文件"""
        test_files = {}
        
        # 创建Word文档
        docx_path = tempfile.mktemp(suffix='.docx')
        doc = Document()
        doc.add_heading('测试文档', 0)
        doc.add_paragraph('这是一个测试段落。')
        doc.add_paragraph('包含多种格式的内容。')
        doc.save(docx_path)
        test_files['docx'] = docx_path
        
        # 创建Markdown文件
        md_path = tempfile.mktemp(suffix='.md')
        with open(md_path, 'w', encoding='utf-8') as f:
            f.write("""# 测试Markdown文档

这是一个**测试**文档，包含：

- 列表项1
- 列表项2
- 列表项3

## 子标题

这里是一些`代码`内容。
""")
        test_files['md'] = md_path
        
        # 创建HTML文件
        html_path = tempfile.mktemp(suffix='.html')
        with open(html_path, 'w', encoding='utf-8') as f:
            f.write("""<!DOCTYPE html>
<html>
<head>
    <title>测试HTML文档</title>
</head>
<body>
    <h1>测试HTML文档</h1>
    <p>这是一个<strong>测试</strong>文档。</p>
    <ul>
        <li>列表项1</li>
        <li>列表项2</li>
    </ul>
</body>
</html>""")
        test_files['html'] = html_path
        
        # 创建Excel文件
        xlsx_path = tempfile.mktemp(suffix='.xlsx')
        df = pd.DataFrame({
            '姓名': ['张三', '李四', '王五'],
            '年龄': [25, 30, 35],
            '城市': ['北京', '上海', '广州']
        })
        df.to_excel(xlsx_path, index=False)
        test_files['xlsx'] = xlsx_path
        
        return test_files
    
    def cleanup_test_files(self, test_files):
        """清理测试文件"""
        for file_path in test_files.values():
            try:
                if os.path.exists(file_path):
                    os.remove(file_path)
            except Exception as e:
                self.log(f"清理文件失败 {file_path}: {e}", "WARNING")
    
    def generate_report(self):
        """生成检查报告"""
        self.log("生成检查报告...")
        
        # 计算总体状态
        all_checks = []
        for category, checks in self.results.items():
            if category == 'overall_status':
                continue
            if isinstance(checks, dict):
                for check_name, check_result in checks.items():
                    if isinstance(check_result, dict) and 'status' in check_result:
                        all_checks.append(check_result['status'])
        
        success_count = all_checks.count('success')
        warning_count = all_checks.count('warning')
        error_count = all_checks.count('error') + all_checks.count('failed')
        
        if error_count == 0 and warning_count == 0:
            self.results['overall_status'] = 'excellent'
        elif error_count == 0:
            self.results['overall_status'] = 'good'
        elif error_count < len(all_checks) / 2:
            self.results['overall_status'] = 'needs_attention'
        else:
            self.results['overall_status'] = 'critical'
        
        # 保存详细报告
        report_path = 'project_check_report.json'
        with open(report_path, 'w', encoding='utf-8') as f:
            json.dump(self.results, f, ensure_ascii=False, indent=2)
        
        # 打印摘要
        print("\n" + "="*60)
        print("📊 项目检查报告摘要")
        print("="*60)
        print(f"✅ 成功: {success_count}")
        print(f"⚠️ 警告: {warning_count}")
        print(f"❌ 错误: {error_count}")
        print(f"📈 总体状态: {self.results['overall_status']}")
        print(f"📄 详细报告: {report_path}")
        print("="*60)
        
        return self.results
    
    def run_full_check(self):
        """运行完整检查"""
        self.log("开始全面项目检查...")
        
        try:
            self.check_dependencies()
            self.check_core_conversion()
            self.check_web_interface()
            self.check_file_support()
            self.check_performance()
            self.check_error_handling()
            
            return self.generate_report()
            
        except Exception as e:
            self.log(f"检查过程中发生异常: {e}", "ERROR")
            traceback.print_exc()
            return None

def main():
    """主函数"""
    print("🔍 DocuFlow 项目全面检查工具")
    print("="*50)
    
    checker = ProjectChecker()
    results = checker.run_full_check()
    
    if results:
        if results['overall_status'] in ['excellent', 'good']:
            print("\n🎉 项目状态良好，所有主要功能正常！")
        elif results['overall_status'] == 'needs_attention':
            print("\n⚠️ 项目需要注意，存在一些问题需要修复。")
        else:
            print("\n🚨 项目存在严重问题，需要立即修复。")
    else:
        print("\n❌ 检查过程失败，请查看错误日志。")

if __name__ == "__main__":
    main()