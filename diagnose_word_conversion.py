#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word转HTML转换问题诊断脚本
帮助用户诊断Word转HTML转换失败的具体原因
"""

import os
import sys
import subprocess
import tempfile
from pathlib import Path
from docx import Document

def check_pandoc_installation():
    """检查Pandoc安装状态"""
    print("\n=== 检查Pandoc安装状态 ===")
    
    try:
        # 检查Pandoc是否安装
        result = subprocess.run(
            ['pandoc', '--version'],
            capture_output=True,
            text=True,
            timeout=10
        )
        
        if result.returncode == 0:
            version_info = result.stdout.split('\n')[0]
            print(f"✓ Pandoc已安装: {version_info}")
            
            # 检查支持的格式
            formats_result = subprocess.run(
                ['pandoc', '--list-input-formats'],
                capture_output=True,
                text=True,
                timeout=10
            )
            
            if 'docx' in formats_result.stdout:
                print("✓ Pandoc支持DOCX格式")
                return True
            else:
                print("❌ Pandoc不支持DOCX格式")
                return False
        else:
            print(f"❌ Pandoc命令执行失败: {result.stderr}")
            return False
            
    except FileNotFoundError:
        print("❌ Pandoc未安装或不在PATH中")
        print("请安装Pandoc: https://pandoc.org/installing.html")
        return False
    except subprocess.TimeoutExpired:
        print("❌ Pandoc命令超时")
        return False
    except Exception as e:
        print(f"❌ 检查Pandoc时发生错误: {str(e)}")
        return False

def check_file_accessibility(file_path):
    """检查文件可访问性"""
    print(f"\n=== 检查文件可访问性 ===")
    print(f"文件路径: {file_path}")
    
    try:
        # 检查文件是否存在
        if not os.path.exists(file_path):
            print("❌ 文件不存在")
            return False
        
        print("✓ 文件存在")
        
        # 检查文件大小
        file_size = os.path.getsize(file_path)
        print(f"✓ 文件大小: {file_size} 字节 ({file_size / 1024 / 1024:.2f} MB)")
        
        if file_size == 0:
            print("❌ 文件为空")
            return False
        
        # 检查文件权限
        if not os.access(file_path, os.R_OK):
            print("❌ 文件无读取权限")
            return False
        
        print("✓ 文件可读取")
        
        # 检查文件扩展名
        file_ext = Path(file_path).suffix.lower()
        if file_ext != '.docx':
            print(f"⚠️  文件扩展名不是.docx: {file_ext}")
        else:
            print("✓ 文件扩展名正确")
        
        return True
        
    except Exception as e:
        print(f"❌ 检查文件时发生错误: {str(e)}")
        return False

def check_docx_file_integrity(file_path):
    """检查DOCX文件完整性"""
    print(f"\n=== 检查DOCX文件完整性 ===")
    
    try:
        # 尝试用python-docx打开文件
        doc = Document(file_path)
        print("✓ 文件可以被python-docx正常打开")
        
        # 检查文档内容
        paragraph_count = len(doc.paragraphs)
        print(f"✓ 文档包含 {paragraph_count} 个段落")
        
        # 检查是否有内容
        has_content = any(p.text.strip() for p in doc.paragraphs)
        if has_content:
            print("✓ 文档包含文本内容")
        else:
            print("⚠️  文档似乎没有文本内容")
        
        return True
        
    except Exception as e:
        print(f"❌ DOCX文件完整性检查失败: {str(e)}")
        print("可能原因:")
        print("  - 文件损坏或不完整")
        print("  - 文件不是标准的DOCX格式")
        print("  - 文件被加密或受保护")
        return False

def test_pandoc_conversion(file_path):
    """测试Pandoc转换"""
    print(f"\n=== 测试Pandoc转换 ===")
    
    try:
        # 创建临时输出文件
        with tempfile.NamedTemporaryFile(suffix='.html', delete=False) as temp_file:
            output_path = temp_file.name
        
        # 构建Pandoc命令
        cmd = [
            'pandoc',
            file_path,
            '-o', output_path,
            '--from', 'docx',
            '--to', 'html',
            '--standalone',
            '--metadata', 'charset=utf-8'
        ]
        
        print(f"执行命令: {' '.join(cmd)}")
        
        # 执行转换
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=60
        )
        
        if result.returncode == 0:
            # 检查输出文件
            if os.path.exists(output_path):
                output_size = os.path.getsize(output_path)
                print(f"✓ 转换成功，输出文件大小: {output_size} 字节")
                
                # 检查输出内容
                with open(output_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                
                if content.strip():
                    print("✓ 输出文件包含内容")
                    print(f"内容预览: {content[:200]}...")
                else:
                    print("⚠️  输出文件为空")
                
                # 清理临时文件
                try:
                    os.unlink(output_path)
                except:
                    pass
                
                return True
            else:
                print("❌ 转换完成但输出文件未生成")
                return False
        else:
            print(f"❌ Pandoc转换失败")
            print(f"错误代码: {result.returncode}")
            print(f"错误信息: {result.stderr}")
            
            # 分析常见错误
            stderr_lower = result.stderr.lower()
            if 'zip' in stderr_lower or 'container' in stderr_lower:
                print("\n可能原因: DOCX文件内部结构损坏")
                print("建议: 尝试用Word重新保存文件")
            elif 'permission' in stderr_lower:
                print("\n可能原因: 文件权限问题")
                print("建议: 检查文件是否被其他程序占用")
            elif 'memory' in stderr_lower:
                print("\n可能原因: 内存不足")
                print("建议: 关闭其他程序或处理较小的文件")
            
            return False
            
    except subprocess.TimeoutExpired:
        print("❌ Pandoc转换超时")
        return False
    except Exception as e:
        print(f"❌ 测试转换时发生错误: {str(e)}")
        return False

def check_file_path_issues(file_path):
    """检查文件路径相关问题"""
    print(f"\n=== 检查文件路径问题 ===")
    
    try:
        # 检查路径长度
        path_length = len(file_path)
        print(f"文件路径长度: {path_length} 字符")
        
        if path_length > 260:
            print("⚠️  文件路径过长，可能在某些系统上造成问题")
        
        # 检查特殊字符
        special_chars = ['"', "'", '&', '<', '>', '|', '?', '*']
        found_special = [char for char in special_chars if char in file_path]
        
        if found_special:
            print(f"⚠️  文件路径包含特殊字符: {found_special}")
            print("建议: 重命名文件，移除特殊字符")
        else:
            print("✓ 文件路径不包含问题字符")
        
        # 检查中文字符
        has_chinese = any('\u4e00' <= char <= '\u9fff' for char in file_path)
        if has_chinese:
            print("⚠️  文件路径包含中文字符，可能在某些环境下造成问题")
            print("建议: 如果转换失败，尝试将文件移动到英文路径")
        else:
            print("✓ 文件路径为英文")
        
        # 检查空格
        if ' ' in file_path:
            print("⚠️  文件路径包含空格")
            print("建议: 如果转换失败，尝试重命名文件移除空格")
        
        return True
        
    except Exception as e:
        print(f"❌ 检查文件路径时发生错误: {str(e)}")
        return False

def provide_solutions():
    """提供解决方案建议"""
    print("\n" + "=" * 60)
    print("🔧 常见解决方案")
    print("=" * 60)
    
    solutions = [
        "1. 文件问题解决方案:",
        "   • 用Microsoft Word打开文件，检查是否正常",
        "   • 如果文件损坏，尝试Word的'打开并修复'功能",
        "   • 将文件另存为新的DOCX文件",
        "",
        "2. 路径问题解决方案:",
        "   • 将文件移动到简短的英文路径",
        "   • 重命名文件，移除特殊字符和空格",
        "   • 避免使用中文文件名和路径",
        "",
        "3. 权限问题解决方案:",
        "   • 确保文件没有被其他程序打开",
        "   • 检查文件和文件夹的读写权限",
        "   • 尝试以管理员身份运行程序",
        "",
        "4. Pandoc问题解决方案:",
        "   • 更新Pandoc到最新版本",
        "   • 重新安装Pandoc",
        "   • 检查Pandoc是否在系统PATH中",
        "",
        "5. 系统资源问题解决方案:",
        "   • 关闭其他占用内存的程序",
        "   • 尝试转换较小的文件",
        "   • 重启计算机释放系统资源"
    ]
    
    for solution in solutions:
        print(solution)

def main():
    """主诊断函数"""
    print("Word转HTML转换问题诊断工具")
    print("=" * 60)
    
    # 获取用户输入的文件路径
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        file_path = input("请输入要诊断的Word文件路径: ").strip()
        if file_path.startswith('"') and file_path.endswith('"'):
            file_path = file_path[1:-1]  # 移除引号
    
    if not file_path:
        print("❌ 未提供文件路径")
        return
    
    print(f"\n诊断文件: {file_path}")
    
    # 执行诊断检查
    checks = [
        ("Pandoc安装状态", lambda: check_pandoc_installation()),
        ("文件可访问性", lambda: check_file_accessibility(file_path)),
        ("DOCX文件完整性", lambda: check_docx_file_integrity(file_path)),
        ("文件路径问题", lambda: check_file_path_issues(file_path)),
        ("Pandoc转换测试", lambda: test_pandoc_conversion(file_path))
    ]
    
    results = []
    for check_name, check_func in checks:
        try:
            result = check_func()
            results.append((check_name, result))
        except Exception as e:
            print(f"❌ {check_name}检查时发生异常: {str(e)}")
            results.append((check_name, False))
    
    # 汇总诊断结果
    print("\n" + "=" * 60)
    print("📊 诊断结果汇总")
    print("=" * 60)
    
    passed = 0
    for check_name, result in results:
        status = "✓ 通过" if result else "❌ 失败"
        print(f"  {check_name}: {status}")
        if result:
            passed += 1
    
    print(f"\n总计: {passed}/{len(results)} 项检查通过")
    
    if passed == len(results):
        print("\n🎉 所有检查都通过了！")
        print("如果转换仍然失败，请检查具体的错误信息。")
    else:
        print("\n⚠️  发现问题，请参考以下解决方案:")
        provide_solutions()

if __name__ == "__main__":
    main()