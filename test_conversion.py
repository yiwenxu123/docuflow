#!/usr/bin/env python3
"""
测试转换功能的脚本
"""

import sys
import os
from pathlib import Path

# 添加项目根目录到Python路径
project_root = Path(__file__).parent
sys.path.insert(0, str(project_root))

from converter.document_converter import DocumentConverter
from utils.config_manager import ConfigManager

def test_pdf_conversion():
    """测试PDF转换"""
    print("=== 测试PDF转换 ===")
    
    # 查找测试PDF文件
    test_files = list(project_root.glob("**/*.pdf"))
    if not test_files:
        print("❌ 未找到PDF测试文件")
        return False
    
    test_file = test_files[0]
    print(f"📄 测试文件: {test_file}")
    
    try:
        # 初始化转换器
        converter = DocumentConverter()
        
        # 执行转换
        output_path = converter.convert_file(
            file_path=str(test_file),
            output_format="md",
            output_dir=str(project_root / "test_output")
        )
        
        if output_path and Path(output_path).exists():
            print(f"✅ PDF转换成功: {output_path}")
            # 显示部分内容
            with open(output_path, 'r', encoding='utf-8') as f:
                content = f.read()[:500]
                print(f"📝 内容预览:\n{content}...")
            return True
        else:
            print("❌ PDF转换失败：输出文件未生成")
            return False
            
    except Exception as e:
        print(f"❌ PDF转换异常: {e}")
        return False

def test_excel_conversion():
    """测试Excel转换"""
    print("\n=== 测试Excel转换 ===")
    
    # 查找测试Excel文件
    test_files = list(project_root.glob("**/*.xlsx")) + list(project_root.glob("**/*.xls"))
    if not test_files:
        print("❌ 未找到Excel测试文件")
        return False
    
    test_file = test_files[0]
    print(f"📊 测试文件: {test_file}")
    
    try:
        # 初始化转换器
        converter = DocumentConverter()
        
        # 执行转换
        output_path = converter.convert_file(
            file_path=str(test_file),
            output_format="md",
            output_dir=str(project_root / "test_output")
        )
        
        if output_path and Path(output_path).exists():
            print(f"✅ Excel转换成功: {output_path}")
            # 显示部分内容
            with open(output_path, 'r', encoding='utf-8') as f:
                content = f.read()[:500]
                print(f"📝 内容预览:\n{content}...")
            return True
        else:
            print("❌ Excel转换失败：输出文件未生成")
            return False
            
    except Exception as e:
        print(f"❌ Excel转换异常: {e}")
        return False

def create_test_files():
    """创建测试文件"""
    print("=== 创建测试文件 ===")
    
    # 创建测试Excel文件
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试数据"
        
        # 添加测试数据
        ws['A1'] = "姓名"
        ws['B1'] = "年龄"
        ws['C1'] = "城市"
        ws['A2'] = "张三"
        ws['B2'] = 25
        ws['C2'] = "北京"
        ws['A3'] = "李四"
        ws['B3'] = 30
        ws['C3'] = "上海"
        
        test_excel = project_root / "test_data.xlsx"
        wb.save(test_excel)
        print(f"✅ 创建测试Excel文件: {test_excel}")
        
    except Exception as e:
        print(f"❌ 创建Excel文件失败: {e}")

def main():
    """主函数"""
    print("🚀 DocuFlow 转换功能测试")
    print("=" * 50)
    
    # 创建输出目录
    output_dir = project_root / "test_output"
    output_dir.mkdir(exist_ok=True)
    
    # 创建测试文件
    create_test_files()
    
    # 测试转换功能
    pdf_result = test_pdf_conversion()
    excel_result = test_excel_conversion()
    
    print("\n" + "=" * 50)
    print("📊 测试结果总结:")
    print(f"PDF转换: {'✅ 成功' if pdf_result else '❌ 失败'}")
    print(f"Excel转换: {'✅ 成功' if excel_result else '❌ 失败'}")
    
    if pdf_result and excel_result:
        print("\n🎉 所有转换测试通过！")
        return True
    else:
        print("\n⚠️  部分转换测试失败，请检查错误信息")
        return False

if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)